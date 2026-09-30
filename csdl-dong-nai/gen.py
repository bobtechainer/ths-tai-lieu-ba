# -*- coding: utf-8 -*-
"""Dựng file Excel rà soát: Sheet Checklist (37 đầu mục) + Sheet Trace chi tiết (từng trường)."""
import io, json, sys
from collections import Counter, OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import build as B
import map2 as M
from build import CO, KHAC, DANG, KHONG

OUT = "Checklist_RaSoat_API_CSDL_DongNai_CThu.xlsx"
API = json.load(io.open("_rs/api_fields.json", encoding="utf-8"))

NHOM = {"truong": range(39, 44), "lop": range(44, 49), "nhansu": range(49, 54),
        "hocsinh": range(54, 59), "pccn": range(59, 64), "ttmh": range(64, 68),
        "lopmon": range(68, 72), "pccm": range(72, 76)}
CUA = {}
for g, rg in NHOM.items():
    for k in rg:
        CUA[k] = g

# khối trong hồ sơ trường không có trên hệ thống
GHI_DT = "Hệ thống không có phần khai báo điểm trường"
GHI_VKK = "Chỉ có ô tích Vùng đặc biệt khó khăn, không khai được danh mục vùng khó khăn"
GHI_HT = "Hệ thống không có phần khai báo hạ tầng Internet, máy tính và thiết bị"
NGOAI_LE_HT = {
    "KINH_TUYEN": (CO, "Trường Học > Kinh tuyến"),
    "VI_TUYEN": (CO, "Trường Học > Vĩ tuyến"),
}


def khoi_truong(path):
    if "DIEM_TRUONG" in path:
        return "DT"
    if "VUNG_KHO_KHAN" in path:
        return "VKK"
    if "HA_TANG_INTERNET" in path:
        return "HT"
    if "THONG_TIN_CAP_NHAT" in path:
        return "TT"
    return ""


thieu_map = []


def tra(stt, path, name):
    """-> (trạng thái, ở đâu / ghi chú, nhóm nội dung, nghĩa tiếng Việt)"""
    g = CUA[stt]
    if g == "truong":
        kh = khoi_truong(path)
        nghia = B.GIAI_TRUONG.get(name, "")
        if kh == "DT":
            return KHONG, GHI_DT, B.NHOM_DT, nghia
        if kh == "VKK":
            return KHONG, GHI_VKK, B.NHOM_DT, nghia
        if kh == "HT":
            if name in NGOAI_LE_HT:
                st, wh = NGOAI_LE_HT[name]
                return st, wh, B.NHOM_HT, nghia
            return KHONG, GHI_HT, B.NHOM_HT, nghia
        v = B.M_TRUONG.get((kh, name))
        if not v:
            thieu_map.append((stt, path, name))
            return "?", "", "", nghia
        return v[0], v[1], v[2], nghia
    if g == "lop":
        v = B.M_LOP.get(name)
        nghia = B.GIAI_LOP.get(name, "")
        if not v:
            thieu_map.append((stt, path, name))
            return "?", "", "", nghia
        return v[0], v[1], v[2], nghia
    if g == "nhansu":
        v = M.M_NS.get(name)
        nghia = M.GIAI_NS.get(name, "")
        if not v:
            thieu_map.append((stt, path, name))
            return "?", "", "", nghia
        return v[0], v[1], v[2], nghia
    if g == "hocsinh":
        v = M.M_HS.get(name)
        nghia = M.GIAI_HS.get(name, "")
        if not v:
            thieu_map.append((stt, path, name))
            return "?", "", "", nghia
        return v[0], v[1], v[2], nghia
    v = M.M_NHO.get(name)
    nghia = M.GIAI_NHO.get(name, "")
    if not v:
        thieu_map.append((stt, path, name))
        return "?", "", "", nghia
    return v[0], v[1], v[2], nghia


# ---------------------------------------------------------------- gom dữ liệu
rows = OrderedDict()
for k in range(39, 76):
    it = API[str(k)]
    out = []
    for f in it["fields"]:
        st, wh, nh, ng = tra(k, f["path"], f["name"])
        out.append({"path": f["path"], "name": f["name"], "bb": f["bb"],
                    "nghia": ng, "st": st, "wh": wh, "nhom": nh})
    rows[k] = {"ten": it["ten"], "fs": out}

if thieu_map:
    print("!! CHUA MAP:")
    for x in thieu_map[:60]:
        print("   ", x)
    sys.exit(1)

# ------------------------------------------------------- kết luận từng đầu mục
CHUA_KL = {67: "Hệ thống chỉ có màn thuộc tính môn học cho tiểu học và THCS/THPT, "
               "chưa tìm thấy màn dành cho GDTX. Cần tài khoản trường GDTX để kiểm lại.",
           71: "Hệ thống chỉ có màn xếp môn học cho tiểu học và THCS/THPT, "
               "chưa tìm thấy màn dành cho GDTX. Cần tài khoản trường GDTX để kiểm lại."}
CAP = {39: "mầm non", 40: "tiểu học", 41: "THCS", 42: "THPT", 43: "GDTX",
       44: "mầm non", 45: "tiểu học", 46: "THCS", 47: "THPT", 48: "GDTX",
       49: "mầm non", 50: "tiểu học", 51: "THCS", 52: "THPT", 53: "GDTX",
       54: "mầm non", 55: "tiểu học", 56: "THCS", 57: "THPT", 58: "GDTX",
       59: "mầm non", 60: "tiểu học", 61: "THCS", 62: "THPT", 63: "GDTX",
       64: "tiểu học", 65: "THCS", 66: "THPT", 67: "GDTX",
       68: "tiểu học", 69: "THCS", 70: "THPT", 71: "GDTX",
       72: "tiểu học", 73: "THCS", 74: "THPT", 75: "GDTX"}


def ghi_chu(k, d):
    fs = d["fs"]
    tong = len(fs)
    co = sum(1 for f in fs if f["st"] in (CO, KHAC))
    dang = sum(1 for f in fs if f["st"] == DANG)
    khong = [f for f in fs if f["st"] == KHONG]
    L = []
    if k in CHUA_KL:
        L.append("Chưa kết luận được.")
        L.append(CHUA_KL[k])
        L.append("Nếu dùng chung màn của THCS/THPT thì các nội dung ngành yêu cầu đều đã có, "
                 "trừ phần môn chuyên đề." if k == 67 else
                 "Nếu dùng chung màn của THCS/THPT thì các nội dung ngành yêu cầu đều đã có.")
        return "\n".join(L)
    if not khong:
        L.append("Hệ thống có đủ %d nội dung ngành yêu cầu." % tong)
        if dang:
            L.append("Trong đó %d nội dung đang lưu ở dạng khác, khi đẩy lên phải quy đổi lại." % dang)
        L.append("Số liệu nhập ở màn: " + ", ".join(sorted(set(
            f["wh"].split(" > ")[0] for f in fs if f["wh"] and "tự xác định" not in f["wh"]))) + ".")
        return "\n".join(L)
    L.append("Ngành yêu cầu %d nội dung, hệ thống có %d, còn thiếu %d."
             % (tong, co + dang, len(khong)))
    if len(khong) <= 3:
        L.append("Còn thiếu: " + "; ".join(f["nghia"] for f in khong) + ".")
    else:
        c = Counter(f["nhom"] for f in khong)
        L.append("Thiếu nhiều nhất ở: " + "; ".join(
            "%s (%d)" % (n, s) for n, s in c.most_common(4)) + ".")
    if dang:
        L.append("Ngoài ra %d nội dung đã có nhưng lưu ở dạng khác, khi đẩy lên phải quy đổi lại." % dang)
    L.append("Xem chi tiết từng nội dung ở sheet Trace chi tiết, lọc theo STT %d." % k)
    return "\n".join(L)


ket = {}
for k, d in rows.items():
    if k in CHUA_KL:
        ket[k] = "Chưa kết luận"
    elif any(f["st"] == KHONG for f in d["fs"]):
        ket[k] = "Không đạt"
    else:
        ket[k] = "Đạt"

# ---------------------------------------------------------------------- Excel
XANH = PatternFill("solid", fgColor="C6EFCE")
DO = PatternFill("solid", fgColor="FFC7CE")
VANG = PatternFill("solid", fgColor="FFEB9C")
XAM = PatternFill("solid", fgColor="F4F6FA")
DAU = PatternFill("solid", fgColor="1F3864")
FDAU = Font(bold=True, color="FFFFFF", size=11)
BD = Border(*[Side(style="thin", color="D0D7E5")] * 4)
TOP = Alignment(vertical="top", wrap_text=True)
MAU_ST = {CO: XANH, KHAC: XANH, DANG: VANG, KHONG: DO}
MAU_KL = {"Đạt": XANH, "Không đạt": DO, "Chưa kết luận": VANG}

wb = Workbook()

# ---- Sheet 1
ws = wb.active
ws.title = "Checklist"
ws.sheet_view.showGridLines = False
H1 = ["STT", "Đầu mục", "Cấp học", "Kết luận", "Ghi chú"]
ws.append(H1)
for i in range(1, len(H1) + 1):
    c = ws.cell(1, i)
    c.fill, c.font, c.border = DAU, FDAU, BD
    c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
ws.row_dimensions[1].height = 30
r = 2
for k, d in rows.items():
    ws.cell(r, 1, k)
    ws.cell(r, 2, d["ten"])
    ws.cell(r, 3, CAP[k])
    ws.cell(r, 4, ket[k])
    ws.cell(r, 5, ghi_chu(k, d))
    for i in range(1, 6):
        c = ws.cell(r, i)
        c.border, c.alignment = BD, TOP
        if r % 2 == 0:
            c.fill = XAM
    ws.cell(r, 4).fill = MAU_KL[ket[k]]
    ws.cell(r, 4).font = Font(bold=True)
    ws.cell(r, 1).alignment = Alignment(vertical="top", horizontal="center")
    ws.cell(r, 3).alignment = Alignment(vertical="top", horizontal="center")
    ws.cell(r, 4).alignment = Alignment(vertical="top", horizontal="center")
    ws.row_dimensions[r].height = 78
    r += 1
for i, w in enumerate([6, 44, 11, 14, 92], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "C2"
ws.auto_filter.ref = "A1:E%d" % (r - 1)

# ---- Sheet 2
w2 = wb.create_sheet("Trace chi tiết")
w2.sheet_view.showGridLines = False
H2 = ["STT", "Đầu mục", "Nhóm nội dung", "Nội dung ngành yêu cầu",
      "Bắt buộc", "Hệ thống có?", "Ở đâu trên hệ thống / vì sao thiếu",
      "Tên trường trong tài liệu API"]
w2.append(H2)
for i in range(1, len(H2) + 1):
    c = w2.cell(1, i)
    c.fill, c.font, c.border = DAU, FDAU, BD
    c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
w2.row_dimensions[1].height = 32
r = 2
for k, d in rows.items():
    for f in d["fs"]:
        w2.cell(r, 1, k)
        w2.cell(r, 2, d["ten"])
        w2.cell(r, 3, f["nhom"])
        w2.cell(r, 4, f["nghia"] or f["name"])
        w2.cell(r, 5, "Bắt buộc" if f["bb"] else "")
        w2.cell(r, 6, f["st"])
        w2.cell(r, 7, f["wh"])
        w2.cell(r, 8, (f["path"] + "." if f["path"] else "") + f["name"])
        for i in range(1, 9):
            c = w2.cell(r, i)
            c.border, c.alignment = BD, TOP
            if k % 2 == 0:
                c.fill = XAM
        w2.cell(r, 6).fill = MAU_ST[f["st"]]
        w2.cell(r, 1).alignment = Alignment(vertical="top", horizontal="center")
        w2.cell(r, 6).alignment = Alignment(vertical="top", horizontal="center")
        r += 1
for i, w in enumerate([6, 38, 34, 46, 10, 17, 56, 52], 1):
    w2.column_dimensions[get_column_letter(i)].width = w
w2.freeze_panes = "D2"
w2.auto_filter.ref = "A1:H%d" % (r - 1)

c0 = Counter(ket.values())
_tf = sum(len(d["fs"]) for d in rows.values())
_sc = sum(1 for d in rows.values() for f in d["fs"] if f["st"] in (CO, KHAC))
_sd = sum(1 for d in rows.values() for f in d["fs"] if f["st"] == DANG)
_sk = sum(1 for d in rows.values() for f in d["fs"] if f["st"] == KHONG)
import huongdan
huongdan.them(wb, c0["Đạt"], c0["Không đạt"], c0["Chưa kết luận"], _tf, _sc, _sd, _sk)

wb.save(OUT)

# ------------------------------------------------------------------- tổng kết
c = Counter(ket.values())
tf = sum(len(d["fs"]) for d in rows.values())
sc = sum(1 for d in rows.values() for f in d["fs"] if f["st"] in (CO, KHAC))
sd = sum(1 for d in rows.values() for f in d["fs"] if f["st"] == DANG)
sk = sum(1 for d in rows.values() for f in d["fs"] if f["st"] == KHONG)
print("Ghi:", OUT)
print("37 dau muc:", dict(c))
print("Trace: %d truong | Co %d | Chua dung dang %d | Khong %d" % (tf, sc, sd, sk))
for k in sorted(rows):
    d = rows[k]
    print("  %2d %-52s %-13s %3d truong, thieu %d" % (
        k, d["ten"][:52], ket[k], len(d["fs"]),
        sum(1 for f in d["fs"] if f["st"] == KHONG)))
