# -*- coding: utf-8 -*-
"""Điền checklist gốc: giữ nguyên format 5 cột, chỉ điền Trạng thái và Ghi chú
cho 37 đầu mục của C Thư (STT 39-75)."""
import io
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core import build_rows
from build import CO, KHAC, DANG, KHONG

OUT = "Checklist_RaSoat_API_CSDL_DongNai_CThu.xlsx"
TIEU_DE = "VI. TIẾP NHẬN XỬ LÝ - CHECKLIST"

# ---- 38 dòng của Em Bách: giữ đúng nguyên trạng trong checklist gốc
BACH = [
 (1, "Tiếp nhận chuyên cần", "Đạt", "cần check lại field MA_BUOI_NGHI vs ĐN"),
 (2, "Tiếp nhận hồ sơ sức khỏe", "Đạt", "Mục 4.7 Nhập Sức khỏe học sinh"),
 (3, "Tiếp nhận đánh giá định kỳ C1", "Chưa kiểm tra", ""),
 (4, "Tiếp nhận đánh giá NLPC C1", "Chưa kiểm tra", ""),
 (5, "Tiếp nhận nhận xét môn học 2C3C", "Chưa kiểm tra", ""),
 (6, "Tiếp nhận nhận xét môn học cấp GDTX", "Chưa kiểm tra", ""),
 (7, "Tiếp nhận nhận xét giáo viên chủ nhiệm", "Chưa kiểm tra", ""),
 (8, "Tiếp nhận nhập điểm C2", "Chưa kiểm tra", ""),
 (9, "Tiếp nhận nhập điểm C3", "Chưa kiểm tra", ""),
 (10, "Tiếp nhận nhập điểm GDTX", "Chưa kiểm tra", ""),
 (11, "Tiếp nhận danh hiệu kết quả thi đua khen thưởng cho THCS, THPT, GDTX", "Chưa kiểm tra", ""),
 (12, "Tiếp nhận hạnh kiểm học sinh cho THCS, THPT, GDTX", "Chưa kiểm tra", ""),
 (13, "Tiếp nhận khen thưởng đột xuất áp dụng cho khối tiểu học", "Chưa kiểm tra", ""),
 (14, "Tiếp nhận khen thưởng cuối năm C1", "Chưa kiểm tra", ""),
 (15, "Tiếp nhận giải thưởng học sinh áp dụng cho THCS, THPT, GDTX", "Chưa kiểm tra", ""),
 (16, "Tiếp nhận hoàn thành chương trình mầm non", "Chưa kiểm tra", ""),
 (17, "Tiếp nhận hoàn thành chương trình tiểu học", "Chưa kiểm tra", ""),
 (18, "Tiếp nhận sức khỏe nuôi dưỡng cấp mầm non", "Chưa kiểm tra", ""),
 (19, "Tiếp nhận tổng kết kết quả đánh giá giáo dục tiểu học", "Chưa kiểm tra", ""),
 (20, "Tiếp nhận tổng kết kết quả học tập THCS", "Chưa kiểm tra", ""),
 (21, "Tiếp nhận tổng kết kết quả học tập THPT", "Chưa kiểm tra", ""),
 (22, "Tiếp nhận tổng kết kết quả học tập GDTX", "Chưa kiểm tra", ""),
 (23, "Tiếp nhận kết quả học tập toàn khóa lớp 12 cấp THPT", "Chưa kiểm tra", ""),
 (24, "Tiếp nhận kết quả học tập toàn khóa lớp 12 cấp GDTX", "Chưa kiểm tra", ""),
 (25, "Tiếp nhận kết quả thể lực học sinh", "Chưa kiểm tra", ""),
 (26, "Tiếp nhận kết quả rèn luyện học sinh cho THCS, THPT, GDTX", "Chưa kiểm tra", ""),
 (27, "Tiếp nhận kết quả xét tốt nghiệp lớp 12 cấp THPT, GDTX", "Chưa kiểm tra", ""),
 (28, "Tiếp nhận kết quả xét tốt nghiệp lớp 9 cấp THCS", "Chưa kiểm tra", ""),
 (29, "Tiếp nhận kết quả xét tốt nghiệp lớp 9 cấp GDTX", "Chưa kiểm tra", ""),
 (30, "Tiếp nhận khóa sổ, hoàn thành tổng kết THCS, THPT, GDTX", "Chưa kiểm tra", ""),
 (31, "Tiếp nhận học sinh đăng ký môn học thi lại cấp THPT", "Chưa kiểm tra", ""),
 (32, "Tiếp nhận học sinh đăng ký môn học thi lại cấp THCS", "Chưa kiểm tra", ""),
 (33, "Tiếp nhận học sinh đăng ký môn học thi lại cấp THPT", "Chưa kiểm tra", ""),
 (34, "Tiếp nhận học sinh đăng ký môn học thi lại cấp GDTX", "Chưa kiểm tra", ""),
 (35, "Tiếp nhận học sinh kết quả thi lại cấp TH", "Chưa kiểm tra", ""),
 (36, "Tiếp nhận học sinh kết quả thi lại cấp THCS", "Chưa kiểm tra", ""),
 (37, "Tiếp nhận học sinh kết quả thi lại cấp THPT", "Chưa kiểm tra", ""),
 (38, "Tiếp nhận học sinh kết quả thi lại cấp GDTX", "Chưa kiểm tra", ""),
]

# ---- tên đầu mục của C Thư: giữ đúng chữ trong checklist gốc
TEN = {
 39: "Tiếp nhận hồ sơ trường cấp mầm non",
 40: "Tiếp nhận hồ sơ trường cấp tiểu học",
 41: "Tiếp nhận hồ sơ trường cấp THCS",
 42: "Tiếp nhận hồ sơ trường cấp THPT",
 43: "Tiếp nhận hồ sơ trường cấp GDTX",
 44: "Tiếp nhận hồ sơ lớp cấp mầm non",
 45: "Tiếp nhận hồ sơ lớp cấp tiểu học",
 46: "Tiếp nhận hồ sơ lớp cấp THCS",
 47: "Tiếp nhận hồ sơ lớp cấp THPT",
 48: "Tiếp nhận hồ sơ lớp cấp GDTX",
 49: "Tiếp nhận hồ sơ nhân sự cấp mầm non",
 50: "Tiếp nhận hồ sơ nhân sự cấp tiểu học",
 51: "Tiếp nhận hồ sơ nhân sự cấp THCS",
 52: "Tiếp nhận hồ sơ nhân sự cấp THPT",
 53: "Tiếp nhận hồ sơ nhân sự cấp GDTX",
 54: "Tiếp nhận hồ sơ học sinh cấp mầm non",
 55: "Tiếp nhận hồ sơ học sinh cấp tiểu học",
 56: "Tiếp nhận hồ sơ học sinh cấp THCS",
 57: "Tiếp nhận hồ sơ học sinh cấp THPT",
 58: "Tiếp nhận hồ sơ học sinh cấp GDTX",
 59: "Tiếp nhận phân công chủ nhiệm lớp mầm non",
 60: "Tiếp nhận phân công chủ nhiệm lớp cấp tiểu học",
 61: "Tiếp nhận phân công chủ nhiệm lớp cấp THCS",
 62: "Tiếp nhận phân công chủ nhiệm lớp cấp THPT",
 63: "Tiếp nhận phân công chủ nhiệm lớp GDTX",
 64: "Tiếp nhận thuộc tính môn học cấp tiểu học",
 65: "Tiếp nhận thuộc tính môn học cấp THCS",
 66: "Tiếp nhận thuộc tính môn học cấp THPT",
 67: "Tiếp nhận thuộc tính môn học GDTX",
 68: "Tiếp nhận thông tin môn học cấp tiểu học",
 69: "Tiếp nhận thông tin lớp môn học cấp THCS",
 70: "Tiếp nhận thông tin lớp môn học cấp THPT",
 71: "Tiếp nhận thông tin lớp môn học GDTX",
 72: "Tiếp nhận phân công chuyên môn nhân sự cấp tiểu học",
 73: "Tiếp nhận phân công chuyên môn nhân sự cấp THCS",
 74: "Tiếp nhận phân công chuyên môn nhân sự cấp THPT",
 75: "Tiếp nhận phân công chuyên môn nhân sự GDTX",
}

CHUA_KL = {
 67: "Chưa kiểm được vì hệ thống chỉ có màn thuộc tính môn học cho tiểu học và cho "
     "THCS/THPT, chưa thấy màn riêng cho GDTX.\n"
     "Nếu GDTX dùng chung màn của THCS/THPT thì chỉ thiếu phần môn chuyên đề.\n"
     "Cần một tài khoản trường GDTX để kết luận.",
 71: "Chưa kiểm được vì hệ thống chỉ có màn xếp môn học cho tiểu học và cho THCS/THPT, "
     "chưa thấy màn riêng cho GDTX.\n"
     "Nếu GDTX dùng chung màn của THCS/THPT thì đã đủ thông tin ngành cần.\n"
     "Cần một tài khoản trường GDTX để kết luận.",
}

rows = build_rows()


def mo_ta_thieu(khong):
    """Mô tả phần còn thiếu bằng lời, gom theo nhóm nội dung."""
    tong_nhom = Counter(f["nhom"] for f in khong)
    ca_phan, rai_rac = [], []
    for nhom, n in tong_nhom.most_common():
        ten = [f["nghia"] for f in khong if f["nhom"] == nhom]
        if n >= 8:
            ca_phan.append("%s (%d thông tin)" % (nhom, n))
        else:
            rai_rac += ten
    L = []
    if ca_phan:
        L.append("Thiếu cả phần: " + "; ".join(ca_phan) + ".")
    if rai_rac:
        if len(rai_rac) > 7:
            L.append("Thiếu rải rác: " + "; ".join(rai_rac[:7]) +
                     " và %d nội dung khác." % (len(rai_rac) - 7))
        else:
            L.append("Còn thiếu: " + "; ".join(rai_rac) + ".")
    return L


def ghi_chu(k):
    if k in CHUA_KL:
        return CHUA_KL[k]
    fs = rows[k]["fs"]
    tong = len(fs)
    dang = [f for f in fs if f["st"] == DANG]
    khong = [f for f in fs if f["st"] == KHONG]
    L = []
    if not khong:
        man = sorted(set(f["wh"].split(" > ")[0] for f in fs
                         if f["wh"] and "tự xác định" not in f["wh"]))
        L.append("Hệ thống có đủ %d thông tin ngành yêu cầu." % tong)
        L.append("Nhập ở màn: " + ", ".join(man) + ".")
        if dang:
            L.append("Có %d thông tin lưu khác dạng ngành cần, khi đẩy lên phải quy đổi lại."
                     % len(dang))
        return "\n".join(L)
    L.append("Ngành cần %d thông tin, hệ thống có %d, còn thiếu %d."
             % (tong, tong - len(khong), len(khong)))
    L += mo_ta_thieu(khong)
    if dang:
        L.append("Ngoài ra %d thông tin đã có nhưng lưu khác dạng ngành cần, "
                 "khi đẩy lên phải quy đổi lại." % len(dang))
    return "\n".join(L)


def trang_thai(k):
    if k in CHUA_KL:
        return "Chưa kết luận"
    return "Không đạt" if any(f["st"] == KHONG for f in rows[k]["fs"]) else "Đạt"


# ---------------------------------------------------------------------- Excel
XANH = PatternFill("solid", fgColor="C6EFCE")
DO = PatternFill("solid", fgColor="FFC7CE")
VANG = PatternFill("solid", fgColor="FFEB9C")
XAM = PatternFill("solid", fgColor="F2F2F2")
DAU = PatternFill("solid", fgColor="D9D9D9")
BD = Border(*[Side(style="thin", color="A6A6A6")] * 4)
TOP = Alignment(vertical="top", wrap_text=True)
GIUA = Alignment(vertical="center", horizontal="center", wrap_text=True)
MAU = {"Đạt": XANH, "Không đạt": DO, "Chưa kết luận": VANG, "Chưa kiểm tra": None}

wb = Workbook()
ws = wb.active
ws.title = "Checklist"

ws.cell(1, 1, TIEU_DE).font = Font(bold=True, size=12)
ws.merge_cells("A1:E1")
ws.cell(1, 1).alignment = Alignment(vertical="center")
ws.row_dimensions[1].height = 22

HD = ["STT", "Đầu mục", "Trạng thái", "Ghi chú", ""]
for i, h in enumerate(HD, 1):
    c = ws.cell(2, i, h)
    c.font, c.fill, c.border, c.alignment = Font(bold=True), DAU, BD, GIUA
ws.row_dimensions[2].height = 24

r = 3
for stt, ten, tt, gc in BACH:
    ws.cell(r, 1, stt); ws.cell(r, 2, ten); ws.cell(r, 3, tt)
    ws.cell(r, 4, gc); ws.cell(r, 5, "Em Bách")
    for i in range(1, 6):
        c = ws.cell(r, i); c.border, c.alignment = BD, TOP
        c.fill = XAM
    ws.cell(r, 1).alignment = GIUA
    ws.cell(r, 3).alignment = GIUA
    ws.cell(r, 5).alignment = GIUA
    if MAU.get(tt):
        ws.cell(r, 3).fill = MAU[tt]
    ws.row_dimensions[r].height = 16
    r += 1

for k in range(39, 76):
    tt, gc = trang_thai(k), ghi_chu(k)
    ws.cell(r, 1, k); ws.cell(r, 2, TEN[k]); ws.cell(r, 3, tt)
    ws.cell(r, 4, gc); ws.cell(r, 5, "C Thư")
    for i in range(1, 6):
        c = ws.cell(r, i); c.border, c.alignment = BD, TOP
    ws.cell(r, 1).alignment = GIUA
    ws.cell(r, 3).alignment = GIUA
    ws.cell(r, 3).font = Font(bold=True)
    ws.cell(r, 3).fill = MAU[tt]
    ws.cell(r, 5).alignment = GIUA
    ws.row_dimensions[r].height = 15 * max(2, gc.count("\n") + 1) + 6
    r += 1

for i, w in enumerate([6, 46, 15, 104, 11], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A3"
ws.auto_filter.ref = "A2:E%d" % (r - 1)
ws.sheet_view.showGridLines = False
wb.save(OUT)

c = Counter(trang_thai(k) for k in range(39, 76))
print("Ghi:", OUT, "|", r - 3, "dong |", dict(c))
for k in range(39, 76):
    print("--- %d %s [%s]" % (k, TEN[k], trang_thai(k)))
    for ln in ghi_chu(k).split("\n"):
        print("    ", ln)
