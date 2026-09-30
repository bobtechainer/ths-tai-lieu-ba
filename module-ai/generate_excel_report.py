from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

out_path = Path(r"c:/Trường học số - source code/bao-cao-module-ai-gkebook-2026-06-03.xlsx")

headers = [
    "STT",
    "Nhóm module",
    "Module (tiếng Việt)",
    "Project",
    "Có AI?",
    "AI đang làm gì?",
    "Có thể gộp?",
    "Mức ưu tiên gộp",
    "Gợi ý AI tiếp theo",
    "Ghi chú / file tham chiếu",
]

rows = [
    [1, "AI dùng chung", "Trợ lý học tập AI / Chatbot toàn cục", "gkebook-school-frontend", "Có", "Chat đa ngữ cảnh theo màn hình; hỗ trợ chế độ tự do, sách số, đọc sách, học theo bài kiểm tra; trả lời dạng stream; gắn nhãn 'Tạo bởi AI'.", "Nên gộp", "Rất cao", "Chuẩn hóa một AI panel chung: hỏi đáp, tóm tắt, giải thích, tạo câu hỏi, tạo ghi chú.", "src/components/Chatbot/GlobalChatbot.tsx; src/services/ai/index.ts; src/hooks/use-chat-with-assistant.tsx"],
    [2, "AI dùng chung", "Trợ lý học tập AI / Chatbot toàn cục", "prj-gkebook-shadcn-frontend", "Có", "Chat đa ngữ cảnh; có chatbot toàn cục và chatbot landing; trả lời dạng stream; gắn nhãn 'Tạo bởi AI'.", "Nên gộp", "Rất cao", "Dùng chung lõi hội thoại, stream parser, answer renderer, quota và context resolver.", "src/components/global-chatbot/index.tsx; src/services/ai/index.ts; src/hooks/use-chat-with-assistant.tsx"],
    [3, "Tài nguyên học tập", "Kho học liệu", "gkebook-school-frontend", "Có gián tiếp", "Có danh sách, chi tiết, chỉnh sửa, duyệt học liệu; AI liên quan qua chatbot, trích từ khóa và xử lý markdown/editor.", "Nên gộp", "Cao", "AI gợi ý học liệu phù hợp theo môn/lớp; AI tóm tắt học liệu; AI sinh câu hỏi từ học liệu.", "src/app/(materials); src/services/ai/index.ts"],
    [4, "Tài nguyên học tập", "Kho học liệu", "prj-gkebook-shadcn-frontend", "Có gián tiếp", "Có danh sách, chi tiết, chỉnh sửa học liệu; nền UI/service rõ để gắn AI vào học liệu.", "Nên gộp", "Cao", "AI tóm tắt học liệu; AI sinh bộ câu hỏi luyện tập nhanh từ học liệu.", "src/app/(materials); src/components/materials"],
    [5, "Tài nguyên học tập", "Bộ sách / Đọc sách", "gkebook-school-frontend", "Có", "Chatbot đổi chế độ theo bối cảnh sách; hỗ trợ hỏi đáp khi đọc sách.", "Nên gộp một phần", "Cao", "AI giải thích đoạn đang đọc; AI tóm tắt chương/bài; AI tạo flashcard từ sách.", "src/app/(book); src/components/Chatbot/GlobalChatbot.tsx"],
    [6, "Tài nguyên học tập", "Bộ sách / Đọc sách", "prj-gkebook-shadcn-frontend", "Có", "Có màn sách, đọc sách, xem trước; chatbot theo ngữ cảnh đọc sách.", "Nên gộp một phần", "Cao", "AI giải thích đoạn bôi đen; AI đọc nội dung thành tiếng; AI tạo bài luyện sau khi đọc.", "src/app/(sach); src/components/global-chatbot/index.tsx"],
    [7, "Kiểm tra đánh giá", "Ngân hàng câu hỏi", "gkebook-school-frontend", "Có một phần", "Có module ngân hàng câu hỏi; AI chủ yếu ở chatbot và phần đặt chỗ tạo đề.", "Nên gộp", "Cao", "AI tìm câu hỏi tương tự; AI gợi ý độ khó; AI sinh giải thích cho câu hỏi.", "src/app/(modules)/ngan-hang-cau-hoi"],
    [8, "Kiểm tra đánh giá", "Ngân hàng câu hỏi", "prj-gkebook-shadcn-frontend", "Có", "Có hook streaming/parse câu hỏi từ tài liệu bằng AI.", "Nên gộp", "Cao", "Dùng chung luồng: đọc tài liệu -> parse câu hỏi -> kiểm duyệt câu hỏi.", "src/hooks/use-question-banks-streamming.tsx"],
    [9, "Học và luyện tập", "Tự học", "gkebook-school-frontend", "Có gián tiếp", "Có module tự học; chatbot có thể hỗ trợ theo ngữ cảnh học/làm bài.", "Có thể gộp", "Trung bình", "AI gợi ý lộ trình học theo kết quả làm bài; AI nhắc phần cần ôn lại.", "src/app/(self-study); src/app/(mission)/tu-hoc"],
    [10, "Học và luyện tập", "Tự học", "prj-gkebook-shadcn-frontend", "Có gián tiếp", "Có module tự học; có nền để gắn trợ lý học tập cá nhân.", "Có thể gộp", "Trung bình", "AI coach theo hồ sơ luyện tập; AI đề xuất 3 bước học tiếp theo.", "src/app/(self-study)/tu-hoc"],
    [11, "Thông tin nội dung", "Tin tức", "gkebook-school-frontend", "Chưa thấy rõ", "Có tin tức/bảng tin nhưng chưa thấy khối AI riêng trong lần rà source.", "Có thể gộp UI", "Thấp", "Có thể mang AI tóm tắt tin tức từ shadcn sang nếu cần.", "src/app/(news)"],
    [12, "Thông tin nội dung", "Tin tức", "prj-gkebook-shadcn-frontend", "Có", "Có phần 'Nội dung do AI tóm tắt' trong chi tiết tin tức.", "Có thể gộp AI helper", "Trung bình", "Chuẩn hóa component tóm tắt AI để dùng cho tin tức, học liệu, sách.", "src/app/(news)/tin-tuc/[slug]/components/NewsSummary.tsx"],
    [13, "Cá nhân hóa", "Thư viện của tôi", "gkebook-school-frontend", "Chưa thấy rõ", "Quản lý học liệu đã tạo, yêu thích, đã tải lên.", "Có thể gộp UI nền", "Trung bình", "AI gợi ý sắp xếp học liệu; AI tìm học liệu liên quan.", "src/app/(my-library)"],
    [14, "Cá nhân hóa", "Thư viện của tôi", "prj-gkebook-shadcn-frontend", "Chưa thấy rõ", "Có thư viện học liệu, thư viện câu hỏi, hồ sơ luyện tập.", "Có thể gộp UI nền", "Trung bình", "AI gợi ý học liệu/câu hỏi dựa trên lịch sử luyện tập.", "src/app/(my-library)"],
    [15, "Kiểm tra đánh giá", "Xưởng tạo đề / Tạo bài kiểm tra", "gkebook-school-frontend", "Đã đặt chỗ AI", "Có mục AI Gen và menu thao tác AI (thêm câu hỏi tương tự/giải thích/tùy chọn) nhưng đang khóa/đang phát triển.", "Nên gộp", "Rất cao", "Hoàn thiện AI tạo đề từ học liệu; cho chọn số câu, mức độ, dạng câu.", "src/mock-data/quiz-studio/quiz-source.tsx; src/components/Quiz/Create/SourceTypes/index.tsx"],
    [16, "Kiểm tra đánh giá", "Xưởng tạo đề / Tạo bài kiểm tra", "prj-gkebook-shadcn-frontend", "Đã đặt chỗ AI", "Có AI Gen trong quiz-types nhưng đang phát triển; có nền parse câu hỏi từ tài liệu.", "Nên gộp", "Rất cao", "Gộp quiz-source, quiz-ai-entry, question parsing và preview flow.", "src/components/quiz-studio/quiz-types/index.tsx; src/mock-data/quiz-studio/quiz-source.tsx"],
    [17, "AI giọng nói", "Chuyển văn bản thành giọng nói", "prj-gkebook-shadcn-frontend", "Có", "Gọi API chuyển chữ thành audio stream.", "Chưa cần gộp ngay", "Thấp", "Đọc tóm tắt, đọc câu hỏi, hỗ trợ học trên di động.", "src/hooks/use-text-to-speech.tsx"],
    [18, "AI kiểm soát sử dụng", "Hạn mức dùng trợ lý AI", "prj-gkebook-shadcn-frontend", "Có", "Theo dõi và hiển thị quota sử dụng chatbot AI.", "Có thể gộp phần lõi", "Trung bình", "Chuẩn hóa quota AI nếu hai project dùng cùng mô hình tính phí/hạn mức.", "src/components/user/user-chatbot-quota; src/constants/index.ts"],
    [19, "AI cho giáo viên", "Đánh giá học liệu của giáo viên bằng AI", "gkebook-school-frontend", "Có", "AI đánh giá mức độ tích cực tạo học liệu của giáo viên ở trang chủ/dashboard.", "Chưa cần gộp ngay", "Thấp", "AI tóm tắt báo cáo giáo viên; AI gợi ý hành động sau khi xem dashboard.", "src/app/(dashboard)/trang-chu/_components/Actions/getAiTeacherMaterialEvaluate.tsx"],
    [20, "Nhà trường", "Lớp học", "gkebook-school-frontend", "Chưa thấy rõ", "Có lớp học, chi tiết lớp, tạo nhiệm vụ, báo cáo học sinh.", "Chưa nên gộp", "Không ưu tiên", "AI tóm tắt tình hình lớp; AI gợi ý giao nhiệm vụ theo nhóm học sinh.", "src/app/(classroom)/lop-hoc"],
    [21, "Nhà trường", "Điểm danh", "gkebook-school-frontend", "Chưa thấy rõ", "Có điểm danh theo lịch/lớp.", "Chưa nên gộp", "Không ưu tiên", "AI phát hiện tín hiệu nghỉ học bất thường; nhắc giáo viên theo dõi thêm.", "src/app/(classroom)/diem-danh"],
    [22, "Nhà trường", "Thời khóa biểu", "gkebook-school-frontend", "Chưa thấy rõ", "Có lịch giảng dạy/thời khóa biểu.", "Chưa nên gộp", "Không ưu tiên", "AI gợi ý chuẩn bị học liệu theo lịch dạy sắp tới.", "src/app/(classroom)/thoi-khoa-bieu"],
    [23, "Nhà trường", "Nhiệm vụ học tập", "gkebook-school-frontend", "Chưa thấy rõ", "Có giao nhiệm vụ, làm bài, chấm điểm.", "Chưa nên gộp toàn bộ", "Thấp", "AI gợi ý nhiệm vụ; AI hỗ trợ chấm tự luận; AI viết nhận xét.", "src/app/(mission)/nhiem-vu"],
    [24, "Kinh doanh", "Gói cước", "prj-gkebook-shadcn-frontend", "Chưa thấy rõ", "Có danh sách gói cước và chi tiết gói.", "Chưa nên gộp", "Không ưu tiên", "AI gợi ý gói phù hợp theo nhu cầu học; làm sau khi có dữ liệu sử dụng.", "src/app/(payment)/goi-cuoc"],
    [25, "Hỗ trợ", "Câu hỏi thường gặp", "prj-gkebook-shadcn-frontend", "Chưa thấy rõ", "Có module hỗ trợ/câu hỏi thường gặp.", "Chưa nên gộp", "Thấp", "AI hỏi đáp theo FAQ; gợi ý bài hướng dẫn liên quan.", "src/app/(support)/cau-hoi-thuong-gap"],
]

wb = Workbook()
ws = wb.active
ws.title = "Bảng tổng hợp"
ws.append(headers)
for row in rows:
    ws.append(row)

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="D9E2F3")

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

widths = [6, 20, 34, 30, 14, 62, 20, 16, 54, 64]
for idx, width in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(idx)].width = width

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
wb.save(out_path)

check_wb = load_workbook(out_path)
check_ws = check_wb.active
print(f"OK: {out_path} rows={check_ws.max_row} cols={check_ws.max_column}")
