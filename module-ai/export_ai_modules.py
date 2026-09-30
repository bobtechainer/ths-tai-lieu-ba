from openpyxl import load_workbook
from pathlib import Path

base = Path(r"c:/Trường học số - source code/_non_code_assets")
source = base / "Danh sách Module hiện có NKG - Đề xuất AI.xlsx"
out = base / "Danh sách Module hiện có NKG - Đề xuất AI - bo-sung-AI.xlsx"

wb = load_workbook(source)

DATA = {
  "Chung": {
    ("1.0", "Kho Học Liệu", ""): ("gkebook-school-frontend, prj-gkebook-shadcn-frontend", "Có", "Có chatbot dùng chung theo ngữ cảnh màn hình; có API trích xuất từ khóa extract-keyword; có luồng chuyển markdown sang Tiptap để hỗ trợ nội dung học liệu.", "Gợi ý học liệu theo hành vi đọc, tóm tắt học liệu, tạo câu hỏi tự luyện từ học liệu"),
    ("2.0", "Quiz", ""): ("gkebook-school-frontend, prj-ths-question-bank", "Có", "Có AI chatbot hỗ trợ ở chế độ học quiz; có check-answer; có trường ai_content_review cho hotspot/nội dung giải thích; có logic template ngữ nghĩa cho nhóm câu hỏi/nhận định.", "Gợi ý đáp án sai thường gặp, sinh lời giải theo mức độ, sinh biến thể câu hỏi cùng chuẩn đầu ra"),
    ("3.0", "Tài liệu số", ""): ("gkebook-school-frontend, prj-gkebook-shadcn-frontend", "Có", "Có chatbot theo ngữ cảnh đọc tài liệu/sách số; có API trích xuất từ khóa; có hỗ trợ chuyển markdown sang định dạng editor.", "Hỏi đáp theo đoạn đang đọc, tóm tắt từng phần, sinh flashcard từ tài liệu"),
    ("4.0", "Bộ sách", ""): ("gkebook-school-frontend, prj-gkebook-shadcn-frontend, prj-ths-smart-class", "Có", "Có chatbot ở chế độ đọc sách số/sách gốc; có shortcut chatbot theo sách; smart class có widget sách để dùng trong lớp.", "Hỏi đáp theo trang sách, sinh câu hỏi kiểm tra nhanh theo bài, gợi ý lộ trình đọc"),
    ("5.0", "Thư viện của tôi", ""): ("gkebook-school-frontend, prj-gkebook-shadcn-frontend", "Chưa thấy rõ", "Đã thấy module thư viện và liên kết với chatbot/material, nhưng chưa thấy luồng AI riêng tách biệt cho thư viện cá nhân.", "Gợi ý học liệu tiếp theo, phân nhóm tự động theo chủ đề, nhắc ôn tập từ tài liệu đã lưu"),
    ("6.0", "Bộ công cụ", ""): ("prj-ths-smart-class", "Có", "Có AutoDraw gọi Google AutoDraw API để nhận gợi ý hình vẽ từ nét bút; có công cụ AI chat trong lớp.", "Nhận diện sơ đồ đang vẽ, gợi ý công cụ phù hợp, chuyển phác thảo thành hình chuẩn"),
    ("7.0", "Tự học", ""): ("gkebook-school-frontend", "Có", "Chatbot được bật theo ngữ cảnh /tu-hoc/; có danh sách câu hỏi gợi ý và hội thoại theo màn hình học.", "Gia sư tự học theo mục tiêu, nhắc ôn theo tiến độ, phát hiện phần học sinh còn nhầm"),
    ("8.0", "AI Tutor", ""): ("gkebook-school-frontend, prj-gkebook-shadcn-frontend, prj-ths-smart-class", "Có", "Có tạo conversation, stream trả lời, sinh tiêu đề hội thoại, quota chatbot, câu hỏi gợi ý, nhiều mode như free, quiz-study, digital-book, original-book; smart class có SidebarAIChat.", "Gia sư theo vai trò giáo viên/học sinh, hỏi đáp bám ngữ cảnh lớp học, gợi ý hoạt động tiếp theo"),
    ("9.0", "Tin tức/Blog", ""): ("gkebook-school-frontend", "Chưa thấy", "Chưa thấy tính năng AI riêng cho tin tức/blog trong phần code đã đọc.", "Tóm tắt bài viết, gợi ý bài liên quan, sinh câu hỏi đọc hiểu từ bài viết"),
    ("10.0", "Ngân hàng câu hỏi", "Ngân hàng"): ("prj-ths-question-bank", "Có", "Có bốc câu hỏi từ ma trận để tạo bộ đề pick-questions; có ai_content_review trong dữ liệu câu hỏi/hotspot; có template nhận định/semantic cho nhóm câu hỏi.", "Gợi ý phân loại câu hỏi, phát hiện câu trùng ý, gợi ý mức độ khó phù hợp"),
    ("11.0", "", "Khung chương trình"): ("prj-ths-question-bank", "Chưa thấy rõ", "Có dữ liệu program/program_id/program_label trong câu hỏi upload và quản lý câu hỏi, nhưng chưa thấy AI riêng cho khung chương trình.", "Gợi ý gán chương trình tự động theo nội dung câu hỏi"),
    ("12.0", "", "Khung độ khó"): ("prj-ths-question-bank", "Chưa thấy rõ", "Có dữ liệu difficulty/difficulty_id/difficulty_label trong luồng upload/preview câu hỏi, nhưng chưa thấy AI chấm độ khó tự động trong phần code đã đọc.", "Gợi ý độ khó tự động, cảnh báo lệch độ khó so với mục tiêu đề"),
    ("13.0", "", "Ma trận đề"): ("prj-ths-question-bank", "Có", "Có dịch vụ sinh suggested_exams từ ma trận bằng API pick-questions; giao diện generate exam hiển thị đề gợi ý để lưu nháp.", "Gợi ý cân bằng ma trận, cảnh báo thiếu chủ đề, sinh nhiều phiên bản đề tương đương"),
    ("14.0", "Smartclass - Lớp học thông minh", ""): ("prj-ths-smart-class", "Có", "Có AI chat trong sidebar lớp học; có AutoDraw; có face engine/embedding để điểm danh khuôn mặt; có analytics useSuggestions; có gợi ý chữ viết tay.", "Tóm tắt buổi học, gợi ý can thiệp theo mức tham gia, hỗ trợ giáo viên trong lúc dạy"),
    ("15.0", "Presentation - Thiết kế bài giảng", ""): ("prj-ths-smart-class, prj-ths-slide-frontend", "Chưa thấy rõ", "Đã thấy module slide/presentation trong smart class, nhưng chưa thấy luồng AI riêng để sinh bài giảng trong phần code đã đọc.", "Sinh dàn ý bài giảng, gợi ý bố cục slide, rút gọn nội dung dài thành slide"),
    ("16.0", "Game mode", ""): ("prj-ths-smart-class, prj-gkebook-frontend-game", "Chưa thấy rõ", "Đã thấy module game/game mode, nhưng chưa thấy AI riêng cho game mode trong phần code đã đọc.", "Sinh câu hỏi nhanh theo chủ đề lớp, điều chỉnh độ khó theo kết quả chơi, gợi ý chế độ chơi phù hợp"),
    ("17.0", "Điểm danh", ""): ("prj-ths-smart-class, gkebook-school-frontend", "Có", "Smart class có FaceProvider, ths-face-engine, upload ảnh khuôn mặt, đăng ký batch embeddings, lấy face embeddings để nhận diện/điểm danh; frontend trường có module điểm danh nhưng chưa thấy AI riêng ở phần đó.", "Cảnh báo điểm danh bất thường, gợi ý xác minh thủ công khi ảnh mờ, thống kê độ ổn định nhận diện"),
    ("18.0", "Authenticator", "Đăng nhập"): ("gkebook-school-frontend, prj-gkebook-shadcn-frontend, prj-ths-smart-class", "Chưa thấy", "Chưa thấy AI riêng cho đăng nhập trong phần code đã đọc.", "Phát hiện đăng nhập bất thường, gợi ý xác thực tăng cường khi có rủi ro"),
    ("19.0", "", "Đăng ký"): ("gkebook-school-frontend, prj-gkebook-shadcn-frontend", "Chưa thấy", "Chưa thấy AI riêng cho đăng ký trong phần code đã đọc.", "Gợi ý điền hồ sơ, phát hiện dữ liệu đăng ký thiếu/không hợp lý"),
    ("20.0", "", "Quên mật khẩu"): ("gkebook-school-frontend, prj-gkebook-shadcn-frontend", "Chưa thấy", "Chưa thấy AI riêng cho quên mật khẩu trong phần code đã đọc.", "Trợ lý hướng dẫn khôi phục tài khoản theo tình huống"),
  },
  "Riêng Nền tảng": {
    ("1.0", "Landing page", ""): ("prj-gkebook-shadcn-frontend, gkebook-school-frontend", "Có", "Có chatbot/assistant ở landing và luồng hội thoại riêng; có hiển thị câu hỏi gợi ý.", "Tư vấn học liệu phù hợp theo nhu cầu ban đầu, gợi ý lộ trình dùng nền tảng"),
    ("2.0", "Gói cước", ""): ("gkebook-school-frontend", "Chưa thấy", "Chưa thấy AI riêng cho gói cước trong phần code đã đọc.", "Gợi ý gói theo nhu cầu sử dụng, ước lượng quota phù hợp"),
    ("1.0", "Quản lý học liệu", ""): ("gkebook-school-frontend, prj-gkebook-shadcn-frontend", "Có", "Có API trích xuất từ khóa, chatbot theo ngữ cảnh học liệu, chuyển markdown sang editor, dữ liệu liên quan material metadata.", "Gắn thẻ tự động, phát hiện học liệu trùng/chưa đủ metadata, gợi ý mô tả học liệu"),
    ("2.0", "Quản lý người dùng", "Phân quyền ngân hàng câu hỏi"): ("prj-ths-question-bank, gkebook-school-frontend", "Chưa thấy", "Chưa thấy AI riêng cho phân quyền trong phần code đã đọc.", "Gợi ý quyền theo vai trò sử dụng thực tế, cảnh báo phân quyền rộng quá mức"),
  },
  "Riêng Site trường": {
    ("1.0", "Lớp học", "Bảng tin"): ("prj-ths-smart-class", "Chưa thấy rõ", "Có AI chat trong lớp học, nhưng chưa thấy AI riêng cho bảng tin lớp trong phần code đã đọc.", "Tóm tắt bảng tin, gợi ý việc cần làm từ thông báo"),
    ("2.0", "", "Báo cáo trong lớp"): ("prj-ths-smart-class", "Có", "Có analytics dashboard, useSuggestions, thẻ insight/assessment trong báo cáo lớp.", "Tóm tắt buổi học theo lớp, gợi ý nhóm học sinh cần hỗ trợ thêm"),
    ("3.0", "", "Học liệu trong lớp"): ("prj-ths-smart-class", "Có", "Có AI chat trong lớp, widget sách/tài liệu, AutoDraw và công cụ hỗ trợ khi dùng học liệu trong phiên học.", "Gợi ý học liệu tiếp theo ngay trong buổi học, hỏi đáp theo tài liệu đang mở"),
    ("4.0", "Nhiệm vụ", ""): ("gkebook-school-frontend", "Có", "Chatbot hoạt động ở ngữ cảnh nhiệm vụ /nhiem-vu/; có AI hỗ trợ trong luồng quiz-study.", "Gợi ý chữa lỗi theo bài làm, nhắc nhiệm vụ chưa hoàn thành theo mức độ ưu tiên"),
    ("5.0", "Lịch giảng dạy", ""): ("prj-ths-smart-class", "Chưa thấy", "Chưa thấy AI riêng cho lịch giảng dạy trong phần code đã đọc.", "Gợi ý phân bổ tiết dạy, nhắc chồng chéo lịch, gợi ý chuẩn bị bài theo lịch"),
    ("6.0", "Báo cáo", ""): ("prj-ths-smart-class", "Có", "Có dashboard phân tích, insight card, suggestion hook, các tab overview/attendance/quiz/student analysis.", "Sinh nhận xét ngắn cho giáo viên, cảnh báo xu hướng giảm tương tác"),
    ("7.0", "Lớp học", "Bảng tin"): ("prj-ths-smart-class", "Chưa thấy rõ", "Có AI chat trong lớp, nhưng chưa thấy AI riêng cho bảng tin phía học sinh trong phần code đã đọc.", "Tóm tắt thông báo quan trọng, nhắc việc cần nộp"),
    ("8.0", "", "Báo cáo trong lớp"): ("prj-ths-smart-class", "Có", "Có báo cáo/analytics cho học sinh trong lớp và theo phiên học.", "Giải thích kết quả học tập cho học sinh bằng ngôn ngữ dễ hiểu"),
    ("9.0", "", "Học liệu trong lớp"): ("prj-ths-smart-class", "Có", "Có AI chat theo ngữ cảnh lớp và tài liệu; có widget học liệu/sách trong lớp.", "Gợi ý phần cần xem lại sau giờ học, tóm tắt nội dung vừa học"),
    ("10.0", "Nhiệm vụ", ""): ("gkebook-school-frontend", "Có", "Có chatbot ở chế độ học nhiệm vụ/quiz-study.", "Gợi ý bước làm tiếp theo, giải thích lỗi sai theo từng câu"),
    ("11.0", "Thời khóa biểu", ""): ("prj-ths-smart-class", "Chưa thấy", "Chưa thấy AI riêng cho thời khóa biểu trong phần code đã đọc.", "Gợi ý lịch học cá nhân, nhắc học theo mức độ tồn đọng"),
    ("12.0", "Báo cáo", ""): ("prj-ths-smart-class", "Có", "Có báo cáo và analytics cho người học, gồm attendance, quiz, engagement, trend.", "Tóm tắt tiến bộ tuần, gợi ý nội dung cần ôn tiếp"),
    ("13.0", "Quản lý học liệu", ""): ("gkebook-school-frontend, prj-gkebook-shadcn-frontend", "Có", "Có AI extract keyword, chatbot ngữ cảnh học liệu, chuyển markdown sang editor.", "Kiểm tra chất lượng học liệu trước khi xuất bản, gợi ý mô tả/ngắn gọn hóa nội dung"),
    ("14.0", "Quản lý lớp học", ""): ("prj-ths-smart-class", "Có", "Có AI chat trong lớp, face attendance, analytics suggestion, AutoDraw trong không gian lớp học.", "Cảnh báo lớp ít tương tác, gợi ý hoạt động lớp phù hợp"),
    ("15.0", "Quản lý người dùng", "Phân quyền ngân hàng câu hỏi"): ("prj-ths-question-bank, prj-ths-shadcn-admin", "Chưa thấy", "Chưa thấy AI riêng cho phân quyền ngân hàng câu hỏi trong phần code đã đọc.", "Gợi ý quyền theo vai trò công việc, cảnh báo tài khoản có quyền bất thường"),
  },
}

NEW_HEADERS = [
  "Project đối chiếu",
  "Tính năng AI hiện có trong code",
  "Mô tả AI hiện có",
  "Gợi ý thêm tính năng AI",
]

for sheet_name, rows in DATA.items():
  ws = wb[sheet_name]
  header_row = None
  for r in range(1, ws.max_row + 1):
    if str(ws.cell(r, 1).value or "").strip() == "STT":
      header_row = r
      break
  if not header_row:
    continue

  start_col = ws.max_column + 1
  for i, header in enumerate(NEW_HEADERS):
    ws.cell(header_row, start_col + i).value = header

  for r in range(header_row + 1, ws.max_row + 1):
    key = (
      "" if ws.cell(r, 1).value is None else str(ws.cell(r, 1).value).strip(),
      "" if ws.cell(r, 2).value is None else str(ws.cell(r, 2).value).strip(),
      "" if ws.cell(r, 3).value is None else str(ws.cell(r, 3).value).strip(),
    )
    if key in rows:
      values = rows[key]
      for i, value in enumerate(values):
        ws.cell(r, start_col + i).value = value

  widths = {0: 30, 1: 24, 2: 70, 3: 60}
  for i, width in widths.items():
    ws.column_dimensions[ws.cell(header_row, start_col + i).column_letter].width = width

wb.save(out)
print(out)
