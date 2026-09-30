from openpyxl import load_workbook
from pathlib import Path

base = Path(r"c:/Trường học số - source code/_non_code_assets")
source = base / "Danh sách Module hiện có NKG - Đề xuất AI.xlsx"
out = base / "Danh sách Module hiện có NKG - Đề xuất AI - humanized.xlsx"

wb = load_workbook(source)

NEW_HEADERS = [
    "Project đối chiếu",
    "Tính năng AI hiện có trong code",
    "Mô tả AI hiện có",
    "Gợi ý thêm tính năng AI",
]

DATA = {
    "Chung": {
        ("1.0", "Kho Học Liệu", ""): ("Trường học số / Nền tảng", "Có", "Có trợ lý để người học hỏi nhanh về nội dung đang xem và được gợi ý học liệu liên quan.", "Tóm tắt học liệu theo đoạn; gợi ý học liệu tiếp theo theo chủ đề; tạo câu hỏi tự luyện từ học liệu."),
        ("2.0", "Quiz", ""): ("Trường học số / Ngân hàng câu hỏi", "Có", "Có trợ lý hỗ trợ khi làm bài: giải thích câu hỏi, gợi ý cách làm và phản hồi theo bài làm.", "Chỉ ra lỗi sai thường gặp; gợi ý bước làm tiếp theo; tạo thêm câu tương tự để luyện ngay."),
        ("3.0", "Tài liệu số", ""): ("Trường học số / Nền tảng", "Có", "Có trợ lý để hỏi đáp khi đọc tài liệu, giúp người học nắm ý chính nhanh hơn.", "Tóm tắt theo mục; tạo flashcard; gợi ý phần nên đọc lại."),
        ("4.0", "Bộ sách", ""): ("Trường học số / Smartclass", "Có", "Có trợ lý hỗ trợ khi đọc sách: hỏi đáp theo nội dung đang học.", "Hỏi nhanh sau mỗi bài; gợi ý lộ trình đọc; nhắc lại khái niệm cốt lõi bằng ví dụ dễ hiểu."),
        ("5.0", "Thư viện của tôi", ""): ("Trường học số / Nền tảng", "Chưa thấy rõ", "Chưa thấy phần trợ lý tách riêng cho thư viện cá nhân; hiện chủ yếu là nơi lưu và quản lý nội dung.", "Tự động nhóm theo chủ đề; nhắc ôn tập theo lịch; gợi ý nội dung mới từ mục đã lưu."),
        ("6.0", "Bộ công cụ", ""): ("Smartclass", "Có", "Có công cụ gợi ý hình vẽ từ nét phác thảo và có trợ lý chat hỗ trợ trong lớp.", "Nhận diện sơ đồ; gợi ý cách trình bày đẹp hơn; chuyển phác thảo thành hình chuẩn để dùng ngay."),
        ("7.0", "Tự học", ""): ("Trường học số", "Có", "Có trợ lý hỗ trợ trong quá trình tự học: hỏi đáp và gợi ý cách học theo nội dung đang làm.", "Nhắc ôn theo tiến độ; gợi ý bài tập phù hợp; tóm tắt phần học hôm nay để tự kiểm tra lại."),
        ("8.0", "AI Tutor", ""): ("Trường học số / Smartclass", "Có", "Có trợ lý trò chuyện để hỏi đáp, gợi ý câu hỏi tiếp theo và hỗ trợ theo từng ngữ cảnh học.", "Thêm chế độ giải thích dễ hiểu; ra bài luyện nhanh; gợi ý câu hỏi để giáo viên dùng kiểm tra miệng."),
        ("9.0", "Tin tức/Blog", ""): ("Trường học số", "Chưa thấy", "Chưa thấy tính năng trợ lý riêng cho phần tin tức và bài viết.", "Tóm tắt bài viết 5 dòng; gợi ý bài liên quan; tạo câu hỏi đọc hiểu ngắn sau khi đọc."),
        ("10.0", "Ngân hàng câu hỏi", "Ngân hàng"): ("Ngân hàng câu hỏi", "Có", "Có phần hỗ trợ tạo và ghép đề từ ma trận; một số câu hỏi đã có phần giải thích hỗ trợ.", "Gợi ý phát hiện câu trùng ý; gợi ý phân loại theo chủ đề; đề xuất mức độ khó phù hợp."),
        ("11.0", "", "Khung chương trình"): ("Ngân hàng câu hỏi", "Chưa thấy rõ", "Có phần gắn chương trình học, nhưng chưa thấy trợ lý hỗ trợ mạnh ở mục này.", "Gợi ý chương trình phù hợp theo nội dung câu hỏi; cảnh báo câu hỏi lệch mạch bài học."),
        ("12.0", "", "Khung dộ khó"): ("Ngân hàng câu hỏi", "Chưa thấy rõ", "Có phần quản lý độ khó, nhưng chưa thấy trợ lý tự động đề xuất rõ ràng.", "Gợi ý độ khó theo dạng câu; cảnh báo đề quá dễ hoặc quá khó so với mục tiêu."),
        ("13.0", "", "Ma trận đề"): ("Ngân hàng câu hỏi", "Có", "Có chức năng tạo đề gợi ý từ ma trận để người dùng chọn và lưu.", "Cảnh báo thiếu chủ đề; gợi ý cân bằng tỉ lệ mức độ; tạo nhiều phiên bản đề tương đương."),
        ("14.0", "Smartclass - Lớp học thông minh", ""): ("Smartclass", "Có", "Trong lớp có trợ lý chat; có công cụ gợi ý hình vẽ; có điểm danh bằng nhận diện khuôn mặt.", "Tóm tắt buổi học cho giáo viên; gợi ý hoạt động tiếp theo khi lớp giảm tương tác; gợi ý hỗ trợ nhóm học sinh cần theo dõi."),
        ("15.0", "Presentation - Thiết kế bài giảng", ""): ("Smartclass / Presentation", "Chưa thấy rõ", "Có phần trình chiếu và bài giảng, nhưng chưa thấy trợ lý tự động tạo nội dung bài giảng.", "Gợi ý dàn ý bài giảng; tự chia nội dung thành slide; gợi ý câu hỏi tương tác chèn vào giữa bài."),
        ("16.0", "Game mode", ""): ("Smartclass / Game", "Chưa thấy rõ", "Có chế độ game, nhưng chưa thấy trợ lý điều chỉnh nội dung game theo lớp học.", "Tự điều chỉnh độ khó theo kết quả chơi; gợi ý câu hỏi nhanh theo chủ đề; đề xuất trò chơi phù hợp theo thời lượng."),
        ("17.0", "Điểm danh", ""): ("Smartclass", "Có", "Có điểm danh bằng nhận diện khuôn mặt trong lớp học thông minh.", "Cảnh báo khi nhận diện chưa chắc; gợi ý kiểm tra lại thủ công; thống kê tỉ lệ nhận diện theo buổi."),
        ("18.0", "Authenticator", "Đăng nhập"): ("Trường học số / Smartclass", "Chưa thấy", "Chưa thấy trợ lý riêng cho phần đăng nhập.", "Cảnh báo đăng nhập bất thường; hướng dẫn xử lý khi có rủi ro; gợi ý tăng bảo mật."),
        ("19.0", "", "Đăng ký"): ("Trường học số", "Chưa thấy", "Chưa thấy trợ lý riêng cho phần đăng ký.", "Hướng dẫn điền thông tin dễ hiểu; gợi ý kiểm tra lại thông tin trước khi gửi."),
        ("20.0", "", "Quên mật khẩu"): ("Trường học số", "Chưa thấy", "Chưa thấy trợ lý riêng cho phần quên mật khẩu.", "Hướng dẫn khôi phục theo từng tình huống; nhắc đổi mật khẩu định kỳ."),
    },
    "Riêng Nền tảng": {
        ("1.0", "Landing page", ""): ("Nền tảng", "Có", "Có trợ lý để người dùng mới hỏi nhanh và được gợi ý nội dung phù hợp.", "Hỏi nhu cầu học; gợi ý lộ trình dùng nhanh; gợi ý học liệu nổi bật theo mục tiêu."),
        ("2.0", "Gói cước", ""): ("Nền tảng", "Chưa thấy", "Chưa thấy trợ lý riêng cho phần gói cước.", "Gợi ý gói theo nhu cầu; ước lượng mức dùng phù hợp; giải thích khác nhau giữa các gói bằng ví dụ."),
        ("1.0", "Quản lý học liệu", ""): ("Nền tảng", "Có", "Có trợ lý hỗ trợ tìm nội dung theo chủ đề và gợi ý nội dung liên quan.", "Gợi ý gắn thẻ; gợi ý mô tả ngắn gọn; cảnh báo nội dung thiếu thông tin quan trọng."),
        ("2.0", "Quản lý người dùng", "Phân quyền ngân hàng câu hỏi"): ("Nền tảng", "Chưa thấy", "Chưa thấy trợ lý riêng cho phần phân quyền.", "Gợi ý quyền theo vai trò; cảnh báo cấp quyền quá rộng; gợi ý mẫu phân quyền theo trường."),
    },
    "Riêng Site trường": {
        ("1.0", "Lớp học", "Bảng tin"): ("Site trường / Smartclass", "Chưa thấy rõ", "Chưa thấy trợ lý riêng cho bảng tin; hiện chủ yếu là nơi xem thông báo.", "Tóm tắt thông báo quan trọng; gợi ý việc cần làm; nhắc hạn nộp theo thông báo."),
        ("2.0", "", "Báo cáo trong lớp"): ("Site trường / Smartclass", "Có", "Có phần gợi ý và nhận xét giúp giáo viên nhìn nhanh tình hình lớp trong buổi học.", "Tóm tắt điểm nổi bật; gợi ý nhóm học sinh cần hỗ trợ; gợi ý câu hỏi kiểm tra nhanh khi lớp giảm tập trung."),
        ("3.0", "", "Học liệu trong lớp"): ("Site trường / Smartclass", "Có", "Có trợ lý hỗ trợ khi dạy trong lớp, gắn với nội dung đang trình chiếu hoặc đang học.", "Gợi ý câu hỏi tương tác; gợi ý hoạt động theo thời lượng; gợi ý học liệu bổ sung cho phần đang dạy."),
        ("4.0", "Nhiệm vụ", ""): ("Site trường", "Có", "Có trợ lý hỗ trợ khi giao và làm nhiệm vụ, gắn với phần học và bài tập.", "Gợi ý chữa lỗi theo bài làm; nhắc nhiệm vụ chưa xong; gợi ý bước làm tiếp theo."),
        ("5.0", "Lịch giảng dạy", ""): ("Site trường", "Chưa thấy", "Chưa thấy trợ lý riêng cho lịch giảng dạy.", "Gợi ý phân bổ tiết dạy; nhắc chồng chéo lịch; gợi ý chuẩn bị bài theo lịch tuần."),
        ("6.0", "Báo cáo", ""): ("Site trường / Smartclass", "Có", "Có phần tổng hợp giúp giáo viên xem nhanh tình hình học tập, điểm danh và mức tham gia.", "Sinh nhận xét ngắn; gợi ý nội dung nên chữa lại; cảnh báo khi mức tham gia giảm dần."),
        ("7.0", "Lớp học", "Bảng tin"): ("Site trường / Smartclass", "Chưa thấy rõ", "Chưa thấy trợ lý riêng cho bảng tin phía học sinh.", "Tóm tắt thông báo quan trọng; nhắc việc cần nộp; gợi ý thứ tự việc nên làm trước."),
        ("8.0", "", "Báo cáo trong lớp"): ("Site trường / Smartclass", "Có", "Có phần giúp học sinh xem lại kết quả học trong lớp và tiến độ của mình.", "Giải thích kết quả bằng lời dễ hiểu; gợi ý phần nên học lại; nhắc mục tiêu cho buổi sau."),
        ("9.0", "", "Học liệu trong lớp"): ("Site trường / Smartclass", "Có", "Có trợ lý hỗ trợ hỏi đáp theo tài liệu đang dùng trong lớp.", "Gợi ý phần cần xem lại sau giờ học; tóm tắt nội dung vừa học; gợi ý tài liệu đọc thêm."),
        ("10.0", "Nhiệm vụ", ""): ("Site trường", "Có", "Có trợ lý hỗ trợ trong lúc làm nhiệm vụ và bài tập.", "Gợi ý bước làm tiếp theo; giải thích lỗi sai theo từng câu; nhắc phần chưa hoàn thành."),
        ("11.0", "Thời khóa biểu", ""): ("Site trường", "Chưa thấy", "Chưa thấy trợ lý riêng cho thời khóa biểu.", "Gợi ý lịch học cá nhân; nhắc chuẩn bị bài; nhắc ôn theo các tiết sắp tới."),
        ("12.0", "Báo cáo", ""): ("Site trường / Smartclass", "Có", "Có phần tổng hợp để học sinh xem tiến độ, mức tham gia và kết quả học.", "Tóm tắt tiến bộ theo tuần; gợi ý nội dung cần ôn tiếp; nhắc mục tiêu học gần nhất."),
        ("13.0", "Quản lý học liệu", ""): ("Site trường / Nền tảng", "Có", "Có hỗ trợ tìm và gợi ý học liệu theo chủ đề.", "Gợi ý mô tả ngắn cho học liệu; gợi ý nhóm nội dung gần nhau; nhắc học liệu còn thiếu thông tin."),
        ("14.0", "Quản lý lớp học", ""): ("Site trường / Smartclass", "Có", "Có trợ lý trong lớp, điểm danh thông minh và phần gợi ý để giáo viên theo dõi lớp học tốt hơn.", "Cảnh báo lớp ít tương tác; gợi ý hoạt động phù hợp; gợi ý nhóm học sinh nên để ý thêm."),
        ("15.0", "Quản lý người dùng", "Phân quyền ngân hàng câu hỏi"): ("Site trường / Quản trị", "Chưa thấy", "Chưa thấy trợ lý riêng cho phần phân quyền người dùng.", "Gợi ý quyền theo vai trò công việc; cảnh báo tài khoản có quyền rộng; gợi ý mẫu phân quyền dùng chung."),
    },
}

for sheet_name, rows in DATA.items():
    ws = wb[sheet_name]
    header_row = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip() == "STT":
            header_row = r
            break
    if not header_row:
        continue

    start_col = ws.max_column + 1
    for i, header in enumerate(NEW_HEADERS):
        ws.cell(header_row, start_col + i).value = header

    for r in range(header_row + 1, ws.max_row + 1):
        key = (
            "" if ws.cell(r, 1).value is None else str(ws.cell(r, 1).value).strip(),
            "" if ws.cell(r, 2).value is None else str(ws.cell(r, 2).value).strip(),
            "" if ws.cell(r, 3).value is None else str(ws.cell(r, 3).value).strip(),
        )
        if key in rows:
            for i, value in enumerate(rows[key]):
                ws.cell(r, start_col + i).value = value

    widths = {0: 28, 1: 18, 2: 55, 3: 50}
    for i, width in widths.items():
        ws.column_dimensions[ws.cell(header_row, start_col + i).column_letter].width = width

wb.save(out)
print(str(out))
