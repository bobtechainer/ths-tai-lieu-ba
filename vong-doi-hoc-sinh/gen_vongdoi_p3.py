# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os

WB_PATH = os.path.join(os.path.dirname(__file__), "vong-doi-hoc-sinh-THS.xlsx")
wb = openpyxl.load_workbook(WB_PATH)

hdr_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
hdr_fill = PatternFill(start_color="2E5A88", end_color="2E5A88", fill_type="solid")
grn_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
org_fill = PatternFill(start_color="E65100", end_color="E65100", fill_type="solid")
body_font = Font(name="Arial", size=10)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))

def hdr(ws, row, n, fill=None):
    for c in range(1, n+1):
        cl = ws.cell(row=row, column=c)
        cl.font = hdr_font; cl.fill = fill or hdr_fill; cl.alignment = wrap; cl.border = thin

def bdy(ws, row, n):
    for c in range(1, n+1):
        cl = ws.cell(row=row, column=c)
        cl.font = body_font; cl.alignment = wrap; cl.border = thin

def tbl(ws, heads, data, sr=1, fill=None):
    for i, h in enumerate(heads, 1):
        ws.cell(row=sr, column=i, value=h)
    hdr(ws, sr, len(heads), fill)
    for r, rd in enumerate(data, sr+1):
        for c, v in enumerate(rd, 1):
            ws.cell(row=r, column=c, value=v)
        bdy(ws, r, len(heads))
    return sr + len(data) + 1

def wcol(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = w

# ==========================================
# SHEET 5: GAP ANALYSIS
# ==========================================
ws5 = wb.create_sheet("5. THS hiện tại và GAP")
wcol(ws5, [5, 24, 38, 40, 14])

tbl(ws5,
    ["STT", "Nghiệp vụ", "THS hiện có", "Cần bổ sung", "Ưu tiên"],
    [
        ["1", "Nhập học",
         "Có form thêm học sinh. Nhập thông tin cá nhân, phụ huynh, gán lớp.",
         "Chưa có quy trình nhập học chính thức (trạng thái 'chờ xếp lớp', 'đã nhập học'). Chưa kiểm tra trùng lặp.", "Cao"],
        ["2", "Xếp lớp",
         "Có gán học sinh vào lớp qua dropdown. Thêm từng em một.",
         "Chưa xếp hàng loạt. Chưa gợi ý cân bằng sĩ số.", "Cao"],
        ["3", "Lên lớp",
         "Chưa có.",
         "Cần tính năng: chọn năm học mới, đánh dấu ai lên lớp / ai đúp, hệ thống tự tạo bản ghi mới ở năm mới. Giữ nguyên dữ liệu năm cũ.", "Rất cao"],
        ["4", "Ở lại lớp (đúp)",
         "Chưa có.",
         "Khi xử lý lên lớp, học sinh đúp ở lại khối cũ. Cần đánh dấu riêng. Thông báo phụ huynh.", "Rất cao"],
        ["5", "Chuyển lớp",
         "Có thể xóa khỏi lớp cũ, thêm vào lớp mới (thủ công, 2 bước).",
         "Cần tính năng 'chuyển lớp' 1 bước: chọn em, chọn lớp đích, xong. Hệ thống tự ghi log.", "Trung bình"],
        ["6", "Chuyển trường (đi)",
         "Chưa có.",
         "Cần: đánh dấu 'đã chuyển trường', xuất hồ sơ số (PDF), cập nhật trạng thái. Học sinh vẫn ở trong hệ thống nhưng bị ẩn.", "Cao"],
        ["7", "Chuyển trường (đến)",
         "Tương đương nhập học mới (không phân biệt).",
         "Cần phân biệt 'nhập học mới' và 'chuyển đến'. Chuyển đến cần nhập trường cũ, lý do.", "Cao"],
        ["8", "Thôi học",
         "Có nút 'vô hiệu hóa' (is_actived = false) nhưng không rõ nghiệp vụ.",
         "Cần trạng thái 'thôi học' riêng. Ghi lý do, ngày thôi học. Lên báo cáo thống kê.", "Cao"],
        ["9", "Bảo lưu",
         "Chưa có.",
         "Cần trạng thái 'bảo lưu'. Ghi thời hạn. Hết hạn thì nhắc trường liên hệ phụ huynh.", "Trung bình"],
        ["10", "Tốt nghiệp",
         "Chưa có.",
         "Cần đánh dấu tốt nghiệp theo đợt. Xuất danh sách. Học sinh tốt nghiệp không hiện ở năm mới.", "Trung bình"],
        ["11", "Tái nhập học",
         "Chưa có.",
         "Tìm lại học sinh cũ (đã thôi / bảo lưu), kích hoạt lại, xếp lớp mới.", "Thấp"],
        ["12", "Lịch sử vòng đời",
         "Chưa có.",
         "Cần bảng ghi lại mọi bước: nhập học, lên lớp, chuyển trường, thôi học... để truy vết toàn bộ quá trình của 1 học sinh.", "Rất cao"],
    ])

# ==========================================
# SHEET 6: ĐỀ XUẤT CHO THS
# ==========================================
ws6 = wb.create_sheet("6. Đề xuất cho THS")
wcol(ws6, [5, 28, 60, 14, 12])

tbl(ws6,
    ["STT", "Đề xuất", "Mô tả chi tiết", "Độ khó", "Ưu tiên"],
    [
        ["1", "Thêm trạng thái cho học sinh",
         "Thêm cột 'trạng thái' vào hồ sơ học sinh. Các giá trị: Đang học, Chuyển trường, Thôi học, Bảo lưu, Tốt nghiệp, Chờ xếp lớp. Cột is_actived hiện tại không đủ để phân biệt.",
         "Thấp", "P0"],
        ["2", "Tạo bảng lịch sử vòng đời",
         "Bảng mới ghi lại toàn bộ lịch sử: em A nhập học ngày X, lên lớp 6 ngày Y, chuyển trường ngày Z... Mỗi dòng là 1 sự kiện. Như một cuốn sổ ghi chép tự động.",
         "Trung bình", "P0"],
        ["3", "Màn hình 'Kết thúc năm học'",
         "Cho phép admin: (1) Chọn năm học hiện tại và năm mới. (2) Đánh dấu từng em: lên lớp, đúp, hoặc tốt nghiệp. (3) Hệ thống tự gán em vào khối mới. (4) Dữ liệu năm cũ giữ nguyên, chỉ tạo thêm ở năm mới.",
         "Cao", "P0"],
        ["4", "Màn hình 'Chuyển trường'",
         "Hai tab: 'Chuyển đi' (chọn em, ghi lý do, xuất hồ sơ PDF, đổi trạng thái) và 'Chuyển đến' (nhập thông tin hoặc import file, ghi trường cũ, xếp lớp).",
         "Trung bình", "P1"],
        ["5", "Màn hình 'Thôi học / Bảo lưu'",
         "Chọn em, chọn lý do (dropdown: hoàn cảnh, sức khỏe, kỷ luật, quá tuổi, khác). Ghi ngày. Cập nhật trạng thái. Bảo lưu: thêm ngày hết hạn.",
         "Thấp", "P1"],
        ["6", "Dashboard thống kê vòng đời",
         "Biểu đồ: sĩ số đầu năm vs cuối năm. Số chuyển đến / chuyển đi / thôi học theo tháng. Tỉ lệ lên lớp / đúp theo khối. So sánh giữa các năm.",
         "Trung bình", "P2"],
        ["7", "Xuất báo cáo theo mẫu Bộ GD",
         "Xuất file Excel/PDF theo mẫu thống kê của Phòng/Sở GD: sĩ số, tăng giảm, phân loại. Để trường đối chiếu với CSDL ngành, không phải nhập lại.",
         "Trung bình", "P2"],
        ["8", "Xếp lớp hàng loạt",
         "Cho phép kéo thả hoặc tự động gán học sinh vào các lớp. Gợi ý cân bằng sĩ số, giới tính. Hiện tại phải thêm từng em một, rất mất thời gian.",
         "Cao", "P1"],
    ])

# ==========================================
# SHEET 7: DATA MODEL GỢI Ý
# ==========================================
ws7 = wb.create_sheet("7. Gợi ý dữ liệu")
wcol(ws7, [28, 18, 55])

r = 1
for c in range(1, 4):
    cl = ws7.cell(row=r, column=c)
    cl.font = hdr_font; cl.fill = grn_fill; cl.border = thin
ws7.cell(row=r, column=1, value="Bổ sung vào hồ sơ học sinh (school_users)")
ws7.cell(row=r, column=2, value=""); ws7.cell(row=r, column=3, value="")
r += 1

heads = ["Tên trường", "Kiểu", "Giải thích"]
data1 = [
    ["status", "Chữ (chọn 1 trong danh sách)",
     "Trạng thái học sinh: Đang học, Chờ xếp lớp, Chuyển trường, Thôi học, Bảo lưu, Tốt nghiệp"],
    ["status_updated_at", "Ngày giờ",
     "Lần cuối trạng thái được thay đổi"],
    ["status_note", "Ghi chú tự do",
     "Lý do khi thay đổi trạng thái (ví dụ: 'Chuyển theo gia đình vào TP.HCM')"],
    ["suspension_end_date", "Ngày",
     "Ngày hết hạn bảo lưu (chỉ dùng khi trạng thái là Bảo lưu)"],
    ["previous_school_name", "Chữ",
     "Tên trường cũ (khi học sinh chuyển đến từ trường khác)"],
]
r = tbl(ws7, heads, data1, r)

r += 1
for c in range(1, 4):
    cl = ws7.cell(row=r, column=c)
    cl.font = hdr_font; cl.fill = org_fill; cl.border = thin
ws7.cell(row=r, column=1, value="Bảng mới: Lịch sử vòng đời (enrollment_history)")
ws7.cell(row=r, column=2, value=""); ws7.cell(row=r, column=3, value="")
r += 1

data2 = [
    ["id", "Mã tự sinh", "Mã riêng cho mỗi dòng lịch sử"],
    ["school_user_id", "Liên kết", "Trỏ đến học sinh nào"],
    ["event_type", "Chữ (chọn 1 trong danh sách)",
     "Loại sự kiện: Nhập học, Lên lớp, Đúp, Chuyển lớp, Chuyển trường đi, Chuyển trường đến, Thôi học, Bảo lưu, Tái nhập, Tốt nghiệp"],
    ["from_class_id", "Liên kết", "Lớp trước đó (trống nếu nhập học lần đầu)"],
    ["to_class_id", "Liên kết", "Lớp mới (trống nếu thôi học hoặc chuyển đi)"],
    ["from_grade_id", "Liên kết", "Khối trước (ví dụ: khối 6)"],
    ["to_grade_id", "Liên kết", "Khối mới (ví dụ: khối 7)"],
    ["from_school_year_id", "Liên kết", "Năm học trước"],
    ["to_school_year_id", "Liên kết", "Năm học mới"],
    ["note", "Ghi chú tự do", "Ghi chú thêm (lý do chuyển, kết quả xét...)"],
    ["effective_date", "Ngày", "Ngày sự kiện có hiệu lực"],
    ["created_by", "Liên kết", "Người thực hiện thao tác"],
    ["created_at", "Ngày giờ", "Thời điểm ghi nhận"],
]
tbl(ws7, heads, data2, r)

wb.save(WB_PATH)
print("Sheet 5-7 OK")
