# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os

WB_PATH = os.path.join(os.path.dirname(__file__), "vong-doi-hoc-sinh-THS.xlsx")
wb = openpyxl.load_workbook(WB_PATH)

hdr_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
hdr_fill = PatternFill(start_color="2E5A88", end_color="2E5A88", fill_type="solid")
grn_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
body_font = Font(name="Arial", size=10)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))

def hdr(ws, row, n, fill=None):
    for c in range(1, n+1):
        cl = ws.cell(row=row, column=c)
        cl.font = hdr_font; cl.fill = fill or hdr_fill; cl.alignment = wrap; cl.border = thin

def bdy(ws, row, n):
    for c in range(1, n+1):
        cl = ws.cell(row=row, column=c)
        cl.font = body_font; cl.alignment = wrap; cl.border = thin

def tbl(ws, heads, data, sr=1, fill=None):
    for i, h in enumerate(heads, 1):
        ws.cell(row=sr, column=i, value=h)
    hdr(ws, sr, len(heads), fill)
    for r, rd in enumerate(data, sr+1):
        for c, v in enumerate(rd, 1):
            ws.cell(row=r, column=c, value=v)
        bdy(ws, r, len(heads))
    return sr + len(data) + 1

def wcol(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = w

# ==========================================
# SHEET 3: CSDL NGÀNH BỘ GIÁO DỤC
# ==========================================
ws3 = wb.create_sheet("3. CSDL ngành Bộ GD")
wcol(ws3, [5, 28, 58, 45])

tbl(ws3,
    ["STT", "Hạng mục", "Mô tả", "Liên hệ với THS"],
    [
        ["1", "Hệ thống CSDL ngành là gì?",
         "Bộ GD&ĐT vận hành hệ thống tại csdl.moet.gov.vn. Mỗi trường có tài khoản riêng để nhập liệu. Hệ thống chia theo cấp: Mầm non, Tiểu học (C1), THCS (C2), THPT (C3), GDTX. Phòng GD, Sở GD, và Bộ GD đều có cổng quản lý riêng để tổng hợp số liệu.",
         "THS cần tương thích với các trường dữ liệu mà CSDL ngành yêu cầu, hoặc hỗ trợ xuất dữ liệu theo format của Bộ để trường không phải nhập 2 lần."],
        ["2", "Thông tin học sinh trên CSDL ngành",
         "Mỗi học sinh có mã định danh riêng. Thông tin gồm: họ tên, ngày sinh, giới tính, dân tộc, địa chỉ, hộ khẩu, đối tượng chính sách (hộ nghèo, con thương binh...), và trạng thái (đang học, đã chuyển, thôi học, tốt nghiệp).",
         "Bảng school_users hiện tại thiếu trường trạng thái (status). Cần bổ sung để đồng bộ với CSDL ngành."],
        ["3", "Lên lớp trên CSDL ngành",
         "Cuối năm học, trường nhập kết quả xếp loại từng học sinh. Hệ thống tự chuyển học sinh sang năm học mới ở khối tiếp theo. Học sinh đúp sẽ ở lại khối cũ. Trường chỉ cần đánh dấu 'lên lớp' hoặc 'ở lại', hệ thống tự xử lý.",
         "THS chưa có tính năng 'kết thúc năm học'. Cần xây dựng quy trình tạo bản ghi mới cho học sinh ở năm học mới, giữ nguyên dữ liệu năm cũ."],
        ["4", "Chuyển trường trên CSDL ngành",
         "Trường cũ 'xuất' học sinh (chuyển trạng thái thành 'đã chuyển'). Trường mới 'nhập' học sinh bằng mã định danh. Dữ liệu liên thông giữa các trường trong cùng tỉnh qua hệ thống của Sở GD. Chuyển khác tỉnh thì phức tạp hơn.",
         "THS hiện không liên thông giữa các trường. Cần thiết kế quy trình 'xuất hồ sơ số' và 'nhập hồ sơ số', hoặc ít nhất là xuất/nhập file."],
        ["5", "Các trạng thái học sinh",
         "CSDL ngành quản lý: Đang học, Chuyển trường, Thôi học, Bảo lưu, Tốt nghiệp, Bị đuổi học. Mỗi lần thay đổi trạng thái đều phải cập nhật trên hệ thống. Nếu không cập nhật, số liệu thống kê sẽ sai.",
         "THS cần thêm trường status vào school_users và ghi log mỗi lần thay đổi trạng thái."],
        ["6", "Báo cáo thống kê",
         "CSDL ngành tự động tổng hợp: sĩ số đầu năm, cuối năm, tăng giảm (chuyển đến, chuyển đi, thôi học, bỏ học...). Các trường phải báo cáo định kỳ cho Phòng/Sở GD. Nếu số liệu trên phần mềm khác với CSDL ngành, trường sẽ gặp khó.",
         "THS nên có dashboard thống kê tương tự để trường đối chiếu. Giảm bớt thời gian trường phải nhập liệu thủ công."],
    ])

# ==========================================
# SHEET 4: BENCHMARK PHẦN MỀM
# ==========================================
ws4 = wb.create_sheet("4. Benchmark phần mềm")
wcol(ws4, [5, 18, 15, 50, 48])

tbl(ws4,
    ["STT", "Phần mềm", "Đơn vị", "Tính năng vòng đời học sinh", "Điểm nổi bật và hạn chế"],
    [
        ["1", "VnEdu", "VNPT",
         "Quản lý hồ sơ học sinh, điểm số, xếp lớp, chuyển trường (liên thông trong tỉnh qua mạng VNPT). Kết nối trực tiếp với CSDL ngành. Xuất báo cáo theo mẫu Bộ GD.",
         "Liên thông tốt trong hệ sinh thái VNPT. Nhiều trường công lập đang dùng. Hạn chế: giao diện cũ, chậm cập nhật tính năng mới."],
        ["2", "SMAS", "Viettel",
         "Tương tự VnEdu: quản lý hồ sơ, điểm, hạnh kiểm, xếp lớp, chuyển trường. Kết nối CSDL ngành. Tích hợp thu học phí qua Viettel Money.",
         "Phổ biến ở miền Bắc. Tích hợp thanh toán tiện lợi. Hạn chế: giao diện cơ bản, ít tùy biến."],
        ["3", "eNetViet", "eNet",
         "Tập trung vào liên lạc nhà trường - phụ huynh (thông báo, nhận xét, bảng điểm). Phần quản lý nghiệp vụ (chuyển trường, lên lớp) không sâu.",
         "Mạnh về giao tiếp và thông báo. Hạn chế: yếu về quản trị nghiệp vụ vòng đời học sinh."],
        ["4", "MISA QLTH", "MISA",
         "Quản lý học sinh, điểm, học bạ, chuyển trường, lên lớp. Xuất báo cáo theo mẫu. Có phiên bản cloud.",
         "Giao diện hiện đại hơn. Tích hợp kế toán trường học. Hạn chế: phí cao hơn, cần đào tạo."],
        ["5", "CSDL ngành (Bộ GD)", "Bộ GD&ĐT",
         "Hệ thống chính thức, bắt buộc mọi trường phải dùng. Quản lý trạng thái học sinh (đang học, chuyển, thôi học, tốt nghiệp). Tổng hợp thống kê toàn ngành.",
         "Bắt buộc sử dụng. Dữ liệu chuyển trường liên thông qua đây. Hạn chế: giao diện web cũ, chậm, hay lỗi vào cao điểm."],
    ])

wb.save(WB_PATH)
print("Sheet 3-4 OK")
