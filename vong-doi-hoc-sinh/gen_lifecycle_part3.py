import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os

WB_PATH = os.path.join(os.path.dirname(__file__), "vong-doi-hoc-sinh-THS.xlsx")
wb = openpyxl.load_workbook(WB_PATH)

hdr_font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
hdr_fill = PatternFill(start_color="2E5A88", end_color="2E5A88", fill_type="solid")
grn_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
body_font = Font(name="Arial", size=10)
wrap = Alignment(wrap_text=True, vertical="top")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

def style_header(ws, row, cols, fill=None):
    for c in range(1, cols+1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font; cell.fill = fill or hdr_fill; cell.alignment = wrap; cell.border = thin_border

def style_row(ws, row, cols):
    for c in range(1, cols+1):
        cell = ws.cell(row=row, column=c)
        cell.font = body_font; cell.alignment = wrap; cell.border = thin_border

def add_table(ws, headers, data, start_row=1, fill=None):
    for i, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=i, value=h)
    style_header(ws, start_row, len(headers), fill)
    for r, row_data in enumerate(data, start_row+1):
        for c, val in enumerate(row_data, 1):
            ws.cell(row=r, column=c, value=val)
        style_row(ws, r, len(headers))
    return start_row + len(data) + 1

# =====================
# SHEET 6: DE XUAT
# =====================
ws6 = wb.create_sheet("6. De xuat cho THS")
ws6.column_dimensions["A"].width = 5
ws6.column_dimensions["B"].width = 25
ws6.column_dimensions["C"].width = 55
ws6.column_dimensions["D"].width = 15
ws6.column_dimensions["E"].width = 15

prop_headers = ["STT", "De xuat", "Mo ta chi tiet", "Do kho", "Uu tien"]
prop_data = [
    ["1", "Them truong status cho hoc sinh",
     "Them cot 'status' vao bang school_users. Cac gia tri: DANG_HOC, CHUYEN_TRUONG, THOI_HOC, BAO_LUU, TOT_NGHIEP, CHO_XEP_LOP. Truong is_actived hien tai khong du de phan biet cac trang thai nay.",
     "Thap", "P0"],
    ["2", "Tao bang enrollment_history",
     "Bang moi ghi lai toan bo lich su: hoc sinh A nhap hoc ngay X, len lop 6 ngay Y, chuyen truong ngay Z... Moi dong la 1 su kien trong vong doi. Cac cot: student_id, event_type, from_class, to_class, from_school_year, to_school_year, note, created_by, created_at.",
     "Trung binh", "P0"],
    ["3", "Man hinh 'Ket thuc nam hoc'",
     "Cho phep admin: (1) Chon nam hoc hien tai va nam hoc moi. (2) Danh dau hoc sinh len lop / dup / tot nghiep. (3) He thong tu tao lop moi (hoac lien ket voi lop da tao) va gan hoc sinh vao. (4) Hoc sinh nam cu van giu nguyen, chi tao ban ghi moi o nam moi.",
     "Cao", "P0"],
    ["4", "Man hinh 'Chuyen truong'",
     "Tab 'Chuyen di': chon hoc sinh, ghi ly do, xuat ho so (PDF). Trang thai hoc sinh chuyen thanh CHUYEN_TRUONG. Tab 'Chuyen den': nhap thong tin hoc sinh moi (hoac import), gan lop, ghi truong cu.",
     "Trung binh", "P1"],
    ["5", "Man hinh 'Thoi hoc / Bao luu'",
     "Chon hoc sinh, chon ly do (dropdown: hoan canh, suc khoe, ky luat, qua tuoi, khac). Ghi ngay hieu luc. Cap nhat trang thai. Voi bao luu: them truong 'ngay het han bao luu'.",
     "Thap", "P1"],
    ["6", "Dashboard thong ke vong doi",
     "Bieu do: so hoc sinh dau nam vs cuoi nam. So chuyen den / chuyen di / thoi hoc / bao luu theo thang. Ti le len lop / dup theo khoi. So sanh giua cac nam hoc.",
     "Trung binh", "P2"],
    ["7", "Xuat bao cao theo mau Bo GD",
     "Xuat file Excel/PDF theo mau thong ke cua Phong/So GD: si so dau nam, cuoi nam, tang giam, phan loai hoc luc, hanh kiem. De truong doi chieu voi CSDL nganh.",
     "Trung binh", "P2"],
    ["8", "Xep lop hang loat",
     "Cho phep keo tha hoac auto-assign hoc sinh vao cac lop. Co goi y can bang si so, gioi tinh, hoc luc. Hien tai phai them tung hoc sinh 1.",
     "Cao", "P1"],
]
add_table(ws6, prop_headers, prop_data)

# =====================
# SHEET 7: DATA MODEL
# =====================
ws7 = wb.create_sheet("7. Data model goi y")
ws7.column_dimensions["A"].width = 25
ws7.column_dimensions["B"].width = 18
ws7.column_dimensions["C"].width = 50

r = 1
ws7.cell(row=r, column=1, value="BANG: school_users (bo sung)")
ws7.cell(row=r, column=1).font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
ws7.cell(row=r, column=1).fill = grn_fill
for c in range(2,4):
    ws7.cell(row=r, column=c).fill = grn_fill
    ws7.cell(row=r, column=c).font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
r += 1

dm1_headers = ["Cot", "Kieu", "Mo ta"]
dm1_data = [
    ["status", "enum/text", "Trang thai hoc sinh: DANG_HOC, CHO_XEP_LOP, CHUYEN_TRUONG, THOI_HOC, BAO_LUU, TOT_NGHIEP, BI_DUOI"],
    ["status_updated_at", "timestamp", "Thoi diem cap nhat trang thai gan nhat"],
    ["status_note", "text", "Ghi chu khi thay doi trang thai (ly do chuyen, thoi hoc...)"],
    ["suspension_end_date", "date", "Ngay het han bao luu (chi dung khi status = BAO_LUU)"],
    ["previous_school_name", "text", "Ten truong cu (khi hoc sinh chuyen den)"],
]
r = add_table(ws7, dm1_headers, dm1_data, r)

r += 1
ws7.cell(row=r, column=1, value="BANG MOI: enrollment_history")
ws7.cell(row=r, column=1).font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
ws7.cell(row=r, column=1).fill = grn_fill
for c in range(2,4):
    ws7.cell(row=r, column=c).fill = grn_fill
    ws7.cell(row=r, column=c).font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
r += 1

dm2_data = [
    ["id", "uuid", "Khoa chinh"],
    ["school_user_id", "uuid (FK)", "Tham chieu den hoc sinh trong school_users"],
    ["event_type", "enum/text", "Loai su kien: NHAP_HOC, LEN_LOP, DUP, CHUYEN_LOP, CHUYEN_TRUONG_DI, CHUYEN_TRUONG_DEN, THOI_HOC, BAO_LUU, TAI_NHAP, TOT_NGHIEP"],
    ["from_class_id", "uuid (FK)", "Lop truoc do (null neu nhap hoc lan dau)"],
    ["to_class_id", "uuid (FK)", "Lop moi (null neu thoi hoc / chuyen truong di)"],
    ["from_school_year_id", "uuid (FK)", "Nam hoc truoc"],
    ["to_school_year_id", "uuid (FK)", "Nam hoc moi"],
    ["from_grade_id", "uuid (FK)", "Khoi truoc"],
    ["to_grade_id", "uuid (FK)", "Khoi moi"],
    ["note", "text", "Ghi chu tu do (ly do chuyen, ket qua xet...)"],
    ["effective_date", "date", "Ngay hieu luc cua su kien"],
    ["created_by", "uuid", "Nguoi thuc hien thao tac"],
    ["created_at", "timestamp", "Thoi diem ghi nhan"],
]
add_table(ws7, dm1_headers, dm2_data, r)

wb.save(WB_PATH)
print("Part 3 saved OK")
