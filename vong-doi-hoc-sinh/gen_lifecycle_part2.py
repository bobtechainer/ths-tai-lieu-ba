import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os

WB_PATH = os.path.join(os.path.dirname(__file__), "vong-doi-hoc-sinh-THS.xlsx")
wb = openpyxl.load_workbook(WB_PATH)

hdr_font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
hdr_fill = PatternFill(start_color="2E5A88", end_color="2E5A88", fill_type="solid")
body_font = Font(name="Arial", size=10)
wrap = Alignment(wrap_text=True, vertical="top")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

def style_header(ws, row, cols):
    for c in range(1, cols+1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = wrap; cell.border = thin_border

def style_row(ws, row, cols, fill=None):
    for c in range(1, cols+1):
        cell = ws.cell(row=row, column=c)
        cell.font = body_font; cell.alignment = wrap; cell.border = thin_border
        if fill: cell.fill = fill

def add_table(ws, headers, data, start_row=1):
    for i, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=i, value=h)
    style_header(ws, start_row, len(headers))
    for r, row_data in enumerate(data, start_row+1):
        for c, val in enumerate(row_data, 1):
            ws.cell(row=r, column=c, value=val)
        style_row(ws, r, len(headers))

# =====================
# SHEET 4: BENCHMARK
# =====================
ws4 = wb.create_sheet("4. Benchmark phan mem")
ws4.column_dimensions["A"].width = 5
ws4.column_dimensions["B"].width = 18
ws4.column_dimensions["C"].width = 15
ws4.column_dimensions["D"].width = 45
ws4.column_dimensions["E"].width = 45

bm_headers = ["STT", "Phan mem", "Don vi", "Tinh nang vong doi hoc sinh", "Diem noi bat"]
bm_data = [
    ["1", "VnEdu", "VNPT",
     "Quan ly ho so hoc sinh, diem so, xep lop, chuyen truong (lien thong trong tinh qua VNPT). Ket noi CSDL nganh. Xuat bao cao theo mau Bo GD.",
     "Lien thong manh trong he sinh thai VNPT. Nhieu truong cong lap dung. Giao dien cu, cham cap nhat."],
    ["2", "SMAS", "Viettel",
     "Tuong tu VnEdu: quan ly ho so, diem, hanh kiem, xep lop, chuyen truong. Tich hop smas.edu.vn. Ket noi CSDL nganh.",
     "Pho bien o mien Bac. Tich hop voi Viettel Money de thu hoc phi. Giao dien co ban."],
    ["3", "eNetViet", "eNet",
     "Tap trung vao lien lac nha truong - phu huynh. Co phan quan ly hoc sinh nhung khong sau ve nghiep vu chuyen truong, len lop.",
     "Manh ve giao tiep (thong bao, nhan xet). Yeu ve quan tri nghiep vu hoc sinh."],
    ["4", "Misa QLTH", "MISA",
     "Phan mem quan ly truong hoc cua MISA. Quan ly hoc sinh, diem, hoc ba, chuyen truong, len lop. Xuat bao cao theo mau.",
     "Giao dien hien dai hon. Co phien ban cloud. Tich hop ke toan."],
    ["5", "MOET CSDL nganh", "Bo GD&DT",
     "He thong chinh thuc. Moi truong bat buoc su dung. Quan ly trang thai hoc sinh (dang hoc, chuyen, thoi hoc, tot nghiep). Tong hop thong ke toan nganh.",
     "Bat buoc. Du lieu chuyen truong lien thong qua he thong nay. Giao dien web co ban, cham."],
]
add_table(ws4, bm_headers, bm_data)

# =====================
# SHEET 5: GAP ANALYSIS
# =====================
ws5 = wb.create_sheet("5. THS hien tai va GAP")
ws5.column_dimensions["A"].width = 5
ws5.column_dimensions["B"].width = 22
ws5.column_dimensions["C"].width = 35
ws5.column_dimensions["D"].width = 35
ws5.column_dimensions["E"].width = 15

gap_headers = ["STT", "Nghiep vu", "THS hien co", "Can bo sung", "Muc uu tien"]
gap_data = [
    ["1", "Nhap hoc",
     "Co form them hoc sinh (add-student-modal). Nhap thong tin ca nhan, phu huynh, gan lop.",
     "Chua co quy trinh nhap hoc chinh thuc (trang thai 'cho xep lop', 'da nhap hoc'). Chua kiem tra trung lap CSDL nganh.",
     "Cao"],
    ["2", "Xep lop",
     "Co gan hoc sinh vao lop qua class_users. Co dropdown chon lop.",
     "Chua co tinh nang xep lop hang loat (batch). Chua co goi y can bang si so.",
     "Cao"],
    ["3", "Len lop",
     "Chua co.",
     "Can tinh nang: chon nam hoc moi, chon hoc sinh len lop, he thong tu tao lop moi (hoac map voi lop cu), copy hoc sinh sang. Giu lai lich su o nam cu.",
     "Rat cao"],
    ["4", "O lai lop (Dup)",
     "Chua co.",
     "Khi len lop, hoc sinh dup se o lai khoi cu. Can danh dau rieng. Can thong bao cho phu huynh.",
     "Rat cao"],
    ["5", "Chuyen lop",
     "Co the xoa khoi lop cu, them vao lop moi (thu cong).",
     "Can tinh nang 'chuyen lop' 1 buoc: chon hoc sinh, chon lop dich, he thong tu cap nhat. Giu log.",
     "Trung binh"],
    ["6", "Chuyen truong (di)",
     "Chua co.",
     "Can tinh nang: danh dau hoc sinh 'da chuyen truong', xuat ho so so (PDF/excel), cap nhat trang thai. Hoc sinh van ton tai trong he thong nhung bi an di.",
     "Cao"],
    ["7", "Chuyen truong (den)",
     "Tuong duong nhap hoc moi.",
     "Can phan biet 'nhap hoc moi' va 'chuyen den'. Chuyen den can nhap thong tin truong cu, ly do chuyen.",
     "Cao"],
    ["8", "Thoi hoc",
     "Co nut 'vo hieu hoa' (is_actived = false) nhung chua ro nghiep vu.",
     "Can trang thai 'thoi hoc' rieng (khac voi 'vo hieu hoa'). Can ghi ly do, ngay thoi hoc. Can bao cao thong ke.",
     "Cao"],
    ["9", "Bao luu",
     "Chua co.",
     "Can trang thai 'bao luu'. Ghi thoi han bao luu. Khi het han, nhac truong lien he phu huynh.",
     "Trung binh"],
    ["10", "Tot nghiep",
     "Chua co.",
     "Can tinh nang danh dau tot nghiep theo dot. Xuat danh sach tot nghiep. Hoc sinh tot nghiep khong con hien thi o nam hoc moi.",
     "Trung binh"],
    ["11", "Tai nhap hoc",
     "Chua co.",
     "Can tim lai hoc sinh cu (da thoi hoc / bao luu), kich hoat lai, xep lop moi.",
     "Thap"],
    ["12", "Lich su vong doi",
     "Chua co.",
     "Can bang enrollment_history ghi lai moi buoc: nhap hoc, len lop, chuyen truong, thoi hoc... de truy vet toan bo qua trinh.",
     "Rat cao"],
]
add_table(ws5, gap_headers, gap_data)

wb.save(WB_PATH)
print("Part 2 saved OK")
