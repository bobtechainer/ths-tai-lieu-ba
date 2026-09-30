# -*- coding: utf-8 -*-
"""Build the English version of the Smartclass 2.0 deck as a .pptx for Canva upload.

Content = the English translation in "Smartclass 2.0 - EN translation.md".
Layout mirrors the original deck (navy / blue / orange, 16:9), text-only so the
slides stay editable after Canva import.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- palette ---
NAVY = RGBColor(0x16, 0x2A, 0x60)
NAVY_D = RGBColor(0x0F, 0x1E, 0x4A)
NAVY_L = RGBColor(0x24, 0x39, 0x77)
BLUE = RGBColor(0x1B, 0x5E, 0xC4)
BLUE_L = RGBColor(0x3E, 0x8E, 0xE0)
TEAL = RGBColor(0x24, 0xB2, 0xC4)
ORANGE = RGBColor(0xF1, 0x59, 0x2A)
ORANGE_L = RGBColor(0xF7, 0xA8, 0x7E)
GREY_BG = RGBColor(0xF4, 0xF5, 0xF7)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xDD, 0xE1, 0xE8)
TEXT = RGBColor(0x3C, 0x43, 0x52)
TEXT_L = RGBColor(0x6B, 0x73, 0x82)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WHITE_D = RGBColor(0xD7, 0xDE, 0xEE)

FONT = "Montserrat"          # Canva ships Montserrat, closest match to the original
SW, SH = 13.333, 7.5
M = 0.7                       # page margin

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers ---
def slide(bg=GREY_BG):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, SW, SH, bg)
    return s


def rect(s, l, t, w, h, fill, line=None, rounded=False, radius=0.06, lw=1):
    shape = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    if rounded:
        shape.adjustments[0] = radius
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(lw)
    shape.shadow.inherit = False
    return shape


def oval(s, l, t, w, h, fill, line=None, lw=1):
    shape = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(lw)
    shape.shadow.inherit = False
    return shape


def text(s, l, t, w, h, blocks, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """blocks: list of dicts -> txt, size, bold, color, italic, space_before, spacing"""
    box = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, b in enumerate(blocks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = b.get("align", align)
        p.line_spacing = b.get("spacing", 1.2)
        if b.get("space_before"):
            p.space_before = Pt(b["space_before"])
        if b.get("space_after"):
            p.space_after = Pt(b["space_after"])
        r = p.add_run()
        r.text = b["txt"]
        f = r.font
        f.name = FONT
        f.size = Pt(b.get("size", 12))
        f.bold = b.get("bold", False)
        f.italic = b.get("italic", False)
        f.color.rgb = b.get("color", TEXT)
    return box


def heading(s, title_txt, sub=None, eyebrow=None, color=NAVY, sub_color=TEXT_L,
            top=0.5, size=30):
    y = top
    if eyebrow:
        text(s, M, y, 10, 0.28, [{"txt": eyebrow, "size": 10.5, "bold": True, "color": BLUE}])
        y += 0.42
    text(s, M, y, SW - 2 * M, 0.62, [{"txt": title_txt, "size": size, "bold": True, "color": color}])
    y += size / 72.0 + 0.16
    if sub:
        text(s, M, y, SW - 2 * M, 0.34,
             [{"txt": sub, "size": 12, "color": sub_color, "italic": True}])
        y += 0.42
    return y


def card(s, l, t, w, h, head, body, accent=NAVY, head_size=13, body_size=10,
         fill=CARD, body_color=TEXT, pad=0.28, rounded=True):
    rect(s, l, t, w, h, fill, line=LINE, rounded=rounded, radius=0.08)
    blocks = [{"txt": head, "size": head_size, "bold": True, "color": accent, "spacing": 1.05}]
    if body:
        blocks.append({"txt": body, "size": body_size, "color": body_color,
                       "space_before": 6, "spacing": 1.25})
    text(s, l + pad, t + pad * 0.8, w - 2 * pad, h - 1.4 * pad, blocks)


def bullet_lines(s, l, t, w, items, size=10.5, color=TEXT, gap=6, bullet="•  "):
    blocks = []
    for i, it in enumerate(items):
        blocks.append({"txt": bullet + it, "size": size, "color": color,
                       "space_before": 0 if i == 0 else gap, "spacing": 1.22})
    text(s, l, t, w, 0.4, blocks)


def kpi_row(s, t, items, h=0.85, l=M, w=SW - 2 * M, accent=NAVY):
    n = len(items)
    gap = 0.16
    cw = (w - gap * (n - 1)) / n
    for i, (val, label) in enumerate(items):
        x = l + i * (cw + gap)
        rect(s, x, t, cw, h, CARD, line=LINE, rounded=True, radius=0.12)
        text(s, x + 0.18, t + 0.12, cw - 0.36, h - 0.24,
             [{"txt": val, "size": 15, "bold": True, "color": accent, "spacing": 1.0},
              {"txt": label, "size": 8.5, "color": TEXT_L, "space_before": 3, "spacing": 1.05}])


# ================================================================ slide 1 ===
s = slide(NAVY)
rect(s, 0, 0, SW, SH, NAVY)
oval(s, M, 3.66, 0.22, 0.22, ORANGE)
text(s, M, 2.35, 11.5, 1.3,
     [{"txt": "Smart School", "size": 60, "bold": True, "color": WHITE, "spacing": 1.0}])
text(s, M + 0.45, 3.58, 10.5, 0.4,
     [{"txt": "An ecosystem for digital transformation in education",
       "size": 17, "color": WHITE_D}])
rect(s, M, 4.55, SW - 2 * M, 0.014, NAVY_L)

# ================================================================ slide 2 ===
s = slide()
text(s, 6.1, 5.35, 6.6, 1.3,
     [{"txt": "CONTENTS", "size": 52, "bold": True, "color": RGBColor(0x1B, 0x1F, 0x27),
       "align": PP_ALIGN.RIGHT}])
oval(s, 7.95, 5.72, 0.46, 0.46, RGBColor(0xC5, 0xD6, 0xF2))
toc = [("I", "The Smart School model at a glance"),
       ("II", "Smart School features in detail"),
       ("III", "Partnership proposal")]
rect(s, 1.15, 1.6, 0.045, 3.55, ORANGE)
for i, (num, label) in enumerate(toc):
    y = 1.45 + i * 1.5
    rect(s, 1.15, y + 0.36, 1.0, 0.045, ORANGE)
    oval(s, 1.02, y + 0.15, 0.44, 0.44, ORANGE)
    oval(s, 1.12, y + 0.25, 0.24, 0.24, GREY_BG)
    oval(s, 2.15, y, 0.95, 0.95, RGBColor(0xE7, 0xE9, 0xEC))
    text(s, 2.15, y + 0.33, 0.95, 0.32,
         [{"txt": num, "size": 13, "bold": True, "color": ORANGE, "align": PP_ALIGN.CENTER}])
    rect(s, 2.9, y + 0.14, 7.6, 0.67, NAVY, rounded=True, radius=0.32)
    text(s, 3.55, y + 0.32, 6.6, 0.34,
         [{"txt": label, "size": 15, "bold": True, "color": WHITE}])

# ================================================================ slide 3 ===
s = slide(WHITE)
heading(s, "Where things stand today",
        "A fragmented experience, with information scattered across systems")
rect(s, 0, 1.95, SW, 1.15, GREY_BG)
bar = ["Time lost, productivity down",
       "A fragmented experience",
       "Reports arrive late and rarely\nanswer what was asked",
       "Digital transformation falls\nshort of expectations"]
for i, b in enumerate(bar):
    x = 0.5 + i * (SW - 1.0) / 4
    text(s, x, 2.2, (SW - 1.0) / 4 - 0.25, 0.7,
         [{"txt": b, "size": 12.5, "bold": True, "color": NAVY, "spacing": 1.25}])
cols = [("1", "In the classroom",
         "Hardware runs on its own, disconnected from the school's management system. "
         "Teachers set everything up by hand at the start of each period, which eats into "
         "class time and invites mistakes."),
        ("2", "At school level",
         "Teachers prepare lessons and enter marks manually. Learning data sits in separate "
         "places, so no one can pull a report quickly when the authorities ask for one."),
        ("3", "At department and ministry level",
         "There is no continuous data flow from the classroom up to the managing agencies. "
         "Reports come late, and there is little real data to judge how digital transformation "
         "is actually going.")]
cw = (SW - 2 * M - 0.5) / 3
for i, (num, head, body) in enumerate(cols):
    x = M + i * (cw + 0.25)
    rect(s, x, 3.6, cw, 2.75, GREY_BG, line=LINE, rounded=True, radius=0.07)
    oval(s, x + 0.3, 3.9, 0.42, 0.42, ORANGE)
    text(s, x + 0.3, 4.0, 0.42, 0.28,
         [{"txt": num, "size": 12, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}])
    text(s, x + 0.3, 4.55, cw - 0.6, 1.6,
         [{"txt": head, "size": 13, "bold": True, "color": NAVY, "spacing": 1.1},
          {"txt": body, "size": 10, "color": TEXT, "space_before": 7, "spacing": 1.28}])

# ================================================================ slide 4 ===
s = slide(WHITE)
heading(s, "What the platform delivers", eyebrow="IGNITE   ·   DIGITAL TRANSFORMATION",
        top=0.42, size=27)
for i, (label, col) in enumerate([("Students", BLUE), ("Teachers", TEAL), ("Schools", ORANGE)]):
    x = M + i * 1.75
    rect(s, x, 1.42, 1.6, 0.46, col, rounded=True, radius=0.4)
    text(s, x, 1.54, 1.6, 0.28,
         [{"txt": label, "size": 11.5, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}])

oval(s, 5.07, 2.25, 3.2, 3.2, NAVY)
text(s, 5.37, 3.25, 2.6, 1.2,
     [{"txt": "NKG", "size": 20, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER,
       "spacing": 1.0},
      {"txt": "SMARTCLASS", "size": 20, "bold": True, "color": WHITE,
       "align": PP_ALIGN.CENTER, "spacing": 1.0},
      {"txt": "The single platform that ties together teaching,\nlearning, management and administration",
       "size": 7.5, "color": WHITE_D, "align": PP_ALIGN.CENTER, "space_before": 8,
       "spacing": 1.15}])

left_items = [("Online learning (in class)", "Live interaction in the classroom, so lessons land better", BLUE),
              ("Personal knowledge library", "Store, organise and share teaching materials and know-how", TEAL),
              ("School operations", "Full oversight of everything the school runs", ORANGE)]
right_items = [("Online learning", "Anytime, anywhere, flexible and effective", BLUE),
               ("Digital resource library", "A rich, standardised library that is easy to draw on", TEAL),
               ("Connected data", "Data synced end to end across the whole system", ORANGE)]
for i, (h, b, c) in enumerate(left_items):
    card(s, M, 2.1 + i * 1.32, 4.0, 1.12, h, b, accent=c, head_size=12, body_size=9.5)
for i, (h, b, c) in enumerate(right_items):
    card(s, SW - M - 4.0, 2.1 + i * 1.32, 4.0, 1.12, h, b, accent=c, head_size=12, body_size=9.5)

rect(s, M, 6.2, SW - 2 * M, 0.95, GREY_BG, line=LINE, rounded=True, radius=0.14)
foot = [("Connected", "Everything and everyone in sync", BLUE),
        ("Effective", "Better teaching, learning and school administration", BLUE),
        ("Secure", "Data protected to international standards", BLUE),
        ("Flexible", "Works anytime, anywhere, on any device", BLUE)]
fw = (SW - 2 * M) / 4
for i, (h, b, c) in enumerate(foot):
    x = M + i * fw
    if i:
        rect(s, x, 6.42, 0.01, 0.5, LINE)
    text(s, x + 0.3, 6.38, fw - 0.55, 0.6,
         [{"txt": h, "size": 11.5, "bold": True, "color": c, "spacing": 1.0},
          {"txt": b, "size": 8.5, "color": TEXT_L, "space_before": 3, "spacing": 1.15}])

# ============================================================== slide 5.1 ===
s = slide(WHITE)
heading(s, "An ecosystem for interactive teaching",
        "AI  ·  Data  ·  Interaction  ·  Personalisation  ·  Assessment",
        eyebrow="IGNITE   ·   DIGITAL TRANSFORMATION", top=0.42, size=27)
text(s, M, 2.05, 6, 0.3, [{"txt": "THE OLD LESSON CYCLE", "size": 10.5, "bold": True, "color": ORANGE}])
cycle = [("1. Lesson prep", "PowerPoint, photocopies, hunting for scattered materials"),
         ("2. Class setup", "Roll call by voice or on paper (3–5 minutes)"),
         ("3. Delivery", "Blackboard and slides, one way only"),
         ("4. Student interaction", "Hands up, spoken answers (only 2–3 students per period)"),
         ("5. Assessment", "Marked by hand, results days later"),
         ("6. Reporting", "Paper mark books, no trend analysis")]
cw = (SW - 2 * M - 0.6) / 6
for i, (h, b) in enumerate(cycle):
    x = M + i * (cw + 0.12)
    card(s, x, 2.45, cw, 1.5, h, b, accent=ORANGE, head_size=10, body_size=8, pad=0.18)

rect(s, M, 4.35, 6.1, 2.55, GREY_BG, line=LINE, rounded=True, radius=0.06)
text(s, M + 0.32, 4.55, 5.5, 0.3, [{"txt": "INPUTS", "size": 10.5, "bold": True, "color": BLUE}])
bullet_lines(s, M + 0.32, 4.95, 5.5, [
    "Teachers — lesson prep, content, learning materials and activities",
    "Students — take part, interact, give feedback",
    "Classroom equipment — AI camera, display, microphone, personal devices",
    "Digital content — video, slides, quizzes, games, reference materials"], size=10)

rect(s, 7.1, 4.35, SW - M - 7.1, 2.55, NAVY, rounded=True, radius=0.06)
text(s, 7.42, 4.55, 5.0, 0.6,
     [{"txt": "AT THE CENTRE", "size": 10.5, "bold": True, "color": ORANGE_L},
      {"txt": "Digital resource and smart interaction platform (AI – Data – Cloud)",
       "size": 11.5, "bold": True, "color": WHITE, "space_before": 4, "spacing": 1.15}])
bullet_lines(s, 7.42, 5.5, 2.5, ["Digital resource library", "AI camera recognition",
                                 "Varied interaction"], size=9.5, color=WHITE_D)
bullet_lines(s, 10.0, 5.5, 2.5, ["Real-time assessment", "AI dashboards and reports",
                                 "Security and privacy"], size=9.5, color=WHITE_D)

# ============================================================== slide 5.2 ===
s = slide(WHITE)
heading(s, "The core toolset, and what it changes",
        eyebrow="AN ECOSYSTEM FOR INTERACTIVE TEACHING", top=0.42, size=27)
text(s, M, 1.75, 6, 0.3, [{"txt": "CORE TOOLSET", "size": 10.5, "bold": True, "color": ORANGE}])
tools = ["Interactive whiteboard + 3D widgets + video",
         "AI attendance camera (2 seconds, fully hands-off)",
         "Real-time quizzes, polls, game mode, chat — the whole class takes part",
         "Automatic marking, running results, personalised feedback",
         "AI dashboard: objectives not yet met, students who need support"]
for i, t_ in enumerate(tools):
    y = 2.15 + i * 0.72
    rect(s, M, y, 6.1, 0.6, GREY_BG, line=LINE, rounded=True, radius=0.14)
    oval(s, M + 0.22, y + 0.19, 0.22, 0.22, BLUE)
    text(s, M + 0.6, y + 0.13, 5.3, 0.4,
         [{"txt": t_, "size": 10, "color": TEXT, "spacing": 1.2}])

rect(s, 7.1, 1.7, SW - M - 7.1, 2.55, RGBColor(0xEF, 0xF6, 0xF0), line=LINE, rounded=True, radius=0.06)
text(s, 7.42, 1.92, 5.0, 0.3, [{"txt": "OUTCOMES", "size": 10.5, "bold": True,
                               "color": RGBColor(0x2E, 0x7D, 0x52)}])
outcomes = [("Higher interaction", "every student takes part; focus and interest go up"),
            ("Deeper understanding", "instant feedback, personalised learning"),
            ("Fairer assessment", "automatic, accurate marking that saves time"),
            ("Decisions based on data", "trend analysis to improve teaching and learning")]
blocks = []
for i, (h, b) in enumerate(outcomes):
    blocks.append({"txt": h, "size": 10.5, "bold": True, "color": NAVY,
                   "space_before": 0 if i == 0 else 8, "spacing": 1.1})
    blocks.append({"txt": b, "size": 9.5, "color": TEXT, "space_before": 1, "spacing": 1.15})
text(s, 7.42, 2.32, 5.0, 1.9, blocks)

rect(s, 7.1, 4.45, SW - M - 7.1, 2.45, NAVY, rounded=True, radius=0.06)
text(s, 7.42, 4.65, 2.4, 0.3, [{"txt": "WHAT THE SCHOOL GAINS", "size": 9.5, "bold": True,
                               "color": ORANGE_L}])
bullet_lines(s, 7.42, 5.0, 2.4, ["Better teaching quality", "Less time spent on operations",
                                 "Standardised processes", "Stronger reputation and brand"],
             size=9, color=WHITE_D, gap=5)
text(s, 10.05, 4.65, 2.5, 0.3, [{"txt": "LONGER TERM", "size": 9.5, "bold": True,
                                "color": ORANGE_L}])
bullet_lines(s, 10.05, 5.0, 2.5, ["Students develop across the board",
                                  "Teachers are more creative and more effective",
                                  "The school's digital transformation is built to last",
                                  "A society where learning runs on smart schools"],
             size=9, color=WHITE_D, gap=5)

# ============================================================== slide 5.3 ===
s = slide(WHITE)
heading(s, "The ecosystem and its partners",
        eyebrow="AN ECOSYSTEM FOR INTERACTIVE TEACHING", top=0.6, size=27)
partners = [("Schools", "Manage and run the system, set the direction for digital transformation and quality", NAVY),
            ("Parents", "Follow their child's progress, stay in step with the school", BLUE),
            ("Content partners", "Quality learning materials, standardised, kept up to date and aligned with the curriculum", TEAL),
            ("Hardware partners", "Devices and technology solutions that are modern, stable and secure", ORANGE)]
pw = (SW - 2 * M - 0.6) / 4
for i, (h, b, c) in enumerate(partners):
    x = M + i * (pw + 0.2)
    rect(s, x, 2.2, pw, 2.9, GREY_BG, line=LINE, rounded=True, radius=0.07)
    rect(s, x, 2.2, pw, 0.09, c, rounded=True, radius=0.4)
    text(s, x + 0.3, 2.6, pw - 0.6, 2.0,
         [{"txt": h, "size": 13.5, "bold": True, "color": c, "spacing": 1.1},
          {"txt": b, "size": 10, "color": TEXT, "space_before": 8, "spacing": 1.3}])
text(s, M, 5.55, SW - 2 * M, 0.5,
     [{"txt": "Data flows one way from the classroom to management, value flows back the other way — "
              "and the ecosystem holds the two together.",
       "size": 11, "italic": True, "color": TEXT_L, "align": PP_ALIGN.CENTER}])

# ================================================================ slide 6 ===
s = slide(NAVY)
text(s, 0, 1.75, SW, 1.0,
     [{"txt": "“", "size": 90, "bold": True, "color": ORANGE, "align": PP_ALIGN.CENTER,
       "spacing": 1.0}])
text(s, 2.1, 3.1, SW - 4.2, 1.8,
     [{"txt": "A smart school is what you get when creative space, technology\n"
              "and educational software come together to make teaching\nand learning genuinely work.",
       "size": 21, "color": WHITE, "align": PP_ALIGN.CENTER, "spacing": 1.5}])

# ================================================================ slide 7 ===
s = slide(WHITE)
heading(s, "What a Smart School makes possible", top=0.55, size=30)
groups = [("Operations",
           ["Full digital transformation: replace scattered teaching practice with a central "
            "LMS/LCMS that keeps 100% of learning materials and progress on one digital platform.",
            "Lower cost over time: 40% less time spent authoring content, since digital materials "
            "can be reused; less spent on printing and traditional facilities."]),
          ("Teaching",
           ["Less administrative load: lesson prep, marking and statistical reporting are automated.",
            "Stronger professional practice: teachers move from \"delivering knowledge\" to guiding "
            "the class, freeing them for creative work and deeper research."]),
          ("The learning experience",
           ["Adaptability: students build the habit of learning on their own and working in teams.",
            "Learning without limits: interactive materials (video, AR/VR, interactive quizzes) "
            "available anytime, anywhere, on a phone.",
            "Better results: an expected 20% gain in how well knowledge sticks, through active "
            "learning and visual interaction."])]
gw = (SW - 2 * M - 0.8) / 3
for i, (h, items) in enumerate(groups):
    x = M + i * (gw + 0.4)
    rect(s, x, 1.95, 0.55, 0.55, NAVY, rounded=True, radius=0.18)
    text(s, x, 2.68, gw, 0.4, [{"txt": h, "size": 17, "bold": True, "color": NAVY}])
    bullet_lines(s, x, 3.25, gw, items, size=10, gap=9)

# ================================================================ slide 8 ===
s = slide(WHITE)
heading(s, "The Smart School ecosystem", top=0.55, size=30)
layers = [("THE SCHOOL", NAVY, 4.6, 1.85,
           ["Education administrators", "Teachers", "Parents", "Students"]),
          ("DIGITAL SCHOOL", BLUE, 5.5, 3.35,
           ["Question bank", "Learning resource library, digital library",
            "Online learning management system (LCMS)"]),
          ("SMART SCHOOL", BLUE_L, 6.4, 4.85,
           ["Smartclass software", "Interactive devices", "AI applications", "End-to-end reporting"])]
for i, (label, col, w, y, items) in enumerate(layers):
    x = M + (6.4 - w) / 2
    rect(s, x, y, w, 1.15, col, rounded=True, radius=0.45)
    text(s, x, y + 0.38, w, 0.4,
         [{"txt": label, "size": 15, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}])
    oval(s, x + w - 0.22, y + 0.42, 0.3, 0.3, WHITE, line=col, lw=1.5)
    text(s, 7.75, y + 0.05, 0.7, 0.3,
         [{"txt": "0%d." % (i + 1), "size": 12, "bold": True, "color": ORANGE}])
    bullet_lines(s, 8.5, y + 0.05, 4.2, items, size=10.5, gap=5)
    rect(s, x + w + 0.05, y + 0.56, 7.6 - (x + w), 0.012, LINE)

# ================================================================ slide 9 ===
s = slide(WHITE)
heading(s, "The LCMS", top=0.55, size=36)
lcms = [("Authoring", "Build lessons directly in the system; no third-party software needed."),
        ("Central repository", "Every asset (images, video, text) kept in one place."),
        ("Reusable content", "Take part of one course and drop it into another."),
        ("Multi-platform delivery",
         "Export to standard formats (SCORM, xAPI, HTML5, PDF) and present straight from the system."),
        ("Personalised content", "Lessons adjust automatically to the learner's profile or ability.")]
for i, (h, b) in enumerate(lcms):
    y = 1.95 + i * 1.02
    rect(s, M, y, 0.08, 0.8, ORANGE if i % 2 == 0 else BLUE)
    text(s, M + 0.4, y + 0.02, 11.4, 0.8,
         [{"txt": h, "size": 14, "bold": True, "color": ORANGE if i % 2 == 0 else BLUE,
           "spacing": 1.05},
          {"txt": b, "size": 11, "color": TEXT, "space_before": 4, "spacing": 1.2}])

# =============================================================== slide 10 ===
s = slide(WHITE)
heading(s, "Information security and\ndocument leak prevention", top=0.5, size=27)
sec = [("Dynamic watermark",
        "Identifying details are overlaid faintly across the screen. If a page is photographed "
        "on the sly, the authorities can trace it straight away."),
       ("Technical barrier against copying",
        "Download, copy/paste and screen capture (PrintScreen) are fully disabled for internal "
        "and classified documents."),
       ("Access log",
        "Every account's access, open and read history is recorded in full, in real time.")]
for i, (h, b) in enumerate(sec):
    y = 2.25 + i * 1.25
    text(s, M, y, 6.0, 1.1,
         [{"txt": h, "size": 13.5, "bold": True, "color": ORANGE, "spacing": 1.05},
          {"txt": b, "size": 10, "color": TEXT, "space_before": 5, "spacing": 1.25}])
rect(s, 7.1, 2.05, SW - M - 7.1, 3.9, GREY_BG, line=LINE, rounded=True, radius=0.06)
text(s, 7.42, 2.3, 5.0, 0.3, [{"txt": "ON SCREEN", "size": 10, "bold": True, "color": NAVY}])
text(s, 7.42, 2.7, 5.0, 1.5,
     [{"txt": "Technical barrier against copying", "size": 11, "bold": True, "color": RGBColor(0xC0, 0x39, 0x2B),
       "spacing": 1.1},
      {"txt": "Download — fully locked (no download button, download API disabled)",
       "size": 9.5, "color": TEXT, "space_before": 6, "spacing": 1.2},
      {"txt": "Copy/Paste — fully blocked (Ctrl+C and right-click disabled)",
       "size": 9.5, "color": TEXT, "space_before": 3, "spacing": 1.2},
      {"txt": "PrintScreen — disabled (the screen capture key is blocked)",
       "size": 9.5, "color": TEXT, "space_before": 3, "spacing": 1.2}])
stats = [("799", "Classified documents\n100% watermarked"), ("127", "Real-time access\nviewing right now"),
         ("100%", "Copy protection"), ("3", "Violations blocked\nthis month")]
for i, (v, l_) in enumerate(stats):
    x = 7.42 + (i % 2) * 2.55
    y = 4.35 + (i // 2) * 0.82
    text(s, x, y, 2.4, 0.7,
         [{"txt": v, "size": 15, "bold": True, "color": NAVY, "spacing": 1.0},
          {"txt": l_, "size": 8, "color": TEXT_L, "space_before": 2, "spacing": 1.15}])
rect(s, 0, 6.35, SW, 0.85, NAVY)
text(s, M, 6.62, SW - 2 * M, 0.4,
     [{"txt": "Meets the strict rules on copying and circulating documents set out in the 2018 Law "
              "on Protection of State Secrets and Decree 26/2020/ND-CP.",
       "size": 11, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}])

# =============================================================== slide 11 ===
s = slide(WHITE)
heading(s, "Training management system",
        "Every data flow automated from the moment information arrives. Manual handling gone entirely.",
        top=0.55, size=30)
tm = [("Automatic data flow",
       "From the moment the list of assigned learners comes in, through to class rosters and timetables."),
      ("Secure examinations",
       "The whole barcode-based anonymisation process for written assignments is digitised. Marks are "
       "stored securely, which holds up as legal grounds for issuing certificates."),
      ("Smart timetable management",
       "The system flags every clash automatically (double-booked rooms, missing equipment, lecturer "
       "schedule conflicts), which makes hall scheduling far easier.")]
cw = (SW - 2 * M - 0.5) / 3
for i, (h, b) in enumerate(tm):
    x = M + i * (cw + 0.25)
    card(s, x, 2.15, cw, 1.95, h, b, accent=NAVY, head_size=13, body_size=10, fill=GREY_BG)
kpi_row(s, 4.45, [("531", "Total learners · 2025–2026 · 21 districts"),
                  ("14", "Classes in session"),
                  ("97.2%", "Attendance rate · up 1.8% on 2024"),
                  ("34", "New guest lecturers"),
                  ("2", "Clashes detected · office to adjust")], h=0.95)
text(s, M, 5.75, SW - 2 * M, 0.9,
     [{"txt": "PROGRAMME MIX", "size": 10, "bold": True, "color": BLUE},
      {"txt": "Intermediate political theory – HC 186 learners (35%)  ·  Advanced political theory 142 (27%)  ·  "
              "State management specialist 108 (20%)  ·  Principal specialist 95 (18%)  ·  6 guest lecturers this week",
       "size": 10, "color": TEXT, "space_before": 5, "spacing": 1.25}])

# =============================================================== slide 12 ===
s = slide(WHITE)
heading(s, "Device management system", top=0.55, size=30)
rect(s, M, 1.55, 0.06, 0.7, ORANGE)
text(s, M + 0.28, 1.6, 11.0, 0.6,
     [{"txt": "If one device goes down, any device with a signal or power fault flashes a red alert "
              "on the administrator's dashboard map straight away.",
       "size": 12, "italic": True, "color": TEXT, "spacing": 1.3}])
dm = [("A clear picture of the estate",
       "One map view shows the status of all 25+ lecture halls as charts and readouts."),
      ("Instant control",
       "Lock an app, check whether a device is online, or push a bulk notification — all from a single point.")]
for i, (h, b) in enumerate(dm):
    card(s, M + i * 3.35, 2.6, 3.15, 1.85, h, b, accent=ORANGE, head_size=12.5, body_size=10,
         fill=GREY_BG)
rect(s, 7.5, 2.6, SW - M - 7.5, 1.85, NAVY, rounded=True, radius=0.07)
text(s, 7.8, 2.8, 4.6, 1.5,
     [{"txt": "ON SCREEN", "size": 9.5, "bold": True, "color": ORANGE_L},
      {"txt": "Device list, room assignment, status, battery and storage, version, last activity",
       "size": 9.5, "color": WHITE_D, "space_before": 6, "spacing": 1.2},
      {"txt": "Quick actions: lock device, sync, screenshot, play sound, wipe data",
       "size": 9.5, "color": WHITE_D, "space_before": 4, "spacing": 1.2},
      {"txt": "App management: choose which apps appear on the student device launcher",
       "size": 9.5, "color": WHITE_D, "space_before": 4, "spacing": 1.2}])
kpi_row(s, 4.85, [("2h 36m", "Average usage per device"), ("886", "Active devices"),
                  ("4,532", "Active users"), ("2,890", "Items blocked")], h=1.0)

# =============================================================== slide 13 ===
s = slide(WHITE)
oval(s, M, 2.35, 0.5, 0.5, ORANGE)
text(s, M, 2.95, 8.0, 2.0,
     [{"txt": "Interactive", "size": 52, "bold": True, "color": ORANGE, "spacing": 1.0},
      {"txt": "whiteboard", "size": 52, "bold": True, "color": ORANGE, "spacing": 1.0}])
text(s, M, 5.15, 8.0, 0.6,
     [{"txt": "Endless space to work in, with a proper set of digital drawing tools.",
       "size": 15, "color": TEXT}])
rect(s, 8.6, 1.5, 4.05, 4.5, GREY_BG, line=LINE, rounded=True, radius=0.05)
text(s, 8.9, 3.45, 3.45, 0.6,
     [{"txt": "Image placeholder\nDrop the screenshot from the original slide here",
       "size": 11, "color": TEXT_L, "align": PP_ALIGN.CENTER, "spacing": 1.35}])

# =============================================================== slide 14 ===
s = slide(NAVY)
text(s, 0, 2.15, SW, 1.0,
     [{"txt": "“", "size": 90, "bold": True, "color": ORANGE, "align": PP_ALIGN.CENTER,
       "spacing": 1.0}])
text(s, 0, 3.55, SW, 0.8,
     [{"txt": "AI in the Smart School", "size": 30, "bold": True, "color": WHITE,
       "align": PP_ALIGN.CENTER}])

# =============================================================== slide 15 ===
s = slide(NAVY)
text(s, M, 0.6, 10, 0.7,
     [{"txt": "AI attendance system", "size": 30, "bold": True, "color": WHITE}])
rect(s, M, 1.6, 7.6, 2.0, NAVY_L, rounded=True, radius=0.06)
text(s, M + 0.35, 1.85, 7.0, 1.6,
     [{"txt": "Tighter discipline", "size": 15, "bold": True, "color": ORANGE, "spacing": 1.05},
      {"txt": "Training regulations are properly enforced, and each class keeps its routine and "
              "discipline. Biometric face recognition puts a stop to people attending or signing in "
              "for someone else — whoever is in the room is the officer or learner who was actually "
              "assigned to the course.",
       "size": 10.5, "color": WHITE_D, "space_before": 7, "spacing": 1.3}])
two = [("Time saved",
        "Attendance happens automatically and quickly, without interrupting the lecturer or the "
        "flow of the lesson."),
       ("Transparent reporting",
        "The system stores attendance data on its own and produces clear statistical reports in "
        "real time.")]
for i, (h, b) in enumerate(two):
    x = M + i * 3.9
    rect(s, x, 3.8, 3.7, 1.85, NAVY_L, rounded=True, radius=0.06)
    text(s, x + 0.32, 4.02, 3.1, 1.5,
         [{"txt": h, "size": 14, "bold": True, "color": ORANGE, "spacing": 1.05},
          {"txt": b, "size": 10, "color": WHITE_D, "space_before": 7, "spacing": 1.3}])
rect(s, 8.65, 1.6, SW - M - 8.65, 4.05, NAVY_D, rounded=True, radius=0.05)
text(s, 8.95, 1.85, 3.4, 0.3, [{"txt": "ON SCREEN", "size": 9.5, "bold": True, "color": ORANGE_L}])
text(s, 8.95, 2.25, 3.4, 3.2,
     [{"txt": "Identity check — random check: look straight into the camera",
       "size": 9.5, "color": WHITE_D, "spacing": 1.25},
      {"txt": "48", "size": 20, "bold": True, "color": WHITE, "space_before": 12, "spacing": 1.0},
      {"txt": "total learners", "size": 8.5, "color": WHITE_D, "spacing": 1.1},
      {"txt": "43  ·  present today (89.6%)", "size": 12, "bold": True, "color": WHITE,
       "space_before": 8, "spacing": 1.1},
      {"txt": "3  ·  absent (2 with permission)", "size": 12, "bold": True, "color": WHITE,
       "space_before": 5, "spacing": 1.1},
      {"txt": "2  ·  biometric alerts (wrong face)", "size": 12, "bold": True, "color": ORANGE,
       "space_before": 5, "spacing": 1.1},
      {"txt": "Attendance sheet — name, position, class, time in, face ID at entry, random check, notes",
       "size": 8.5, "color": WHITE_D, "space_before": 10, "spacing": 1.2}])

# =============================================================== slide 16 ===
s = slide(NAVY)
text(s, M, 0.6, 10, 0.7,
     [{"txt": "AI support for research", "size": 30, "bold": True, "color": WHITE}])
rect(s, M, 1.6, 0.06, 1.05, ORANGE)
text(s, M + 0.3, 1.6, 11.4, 1.0,
     [{"txt": "The AI is trained entirely in-house on the school's own data. Question-and-answer "
              "content is never sent out to public servers (no reliance on OpenAI/ChatGPT), which "
              "keeps information secure and ideologically sound.",
       "size": 12, "italic": True, "color": WHITE_D, "spacing": 1.35}])
res = [("Data analysis support",
        "Summarises and visualises figures; processes data to whatever output is required."),
       ("Suggesting an approach",
        "Researches, analyses and proposes a suitable framework and methodology."),
       ("Challenging the argument",
        "Asks probing questions back to surface unstated assumptions, so an argument is properly "
        "stress-tested before a draft, a research paper or an oral defence goes forward.")]
cw = (SW - 2 * M - 0.6) / 3
for i, (h, b) in enumerate(res):
    x = M + i * (cw + 0.3)
    rect(s, x, 3.1, cw, 2.65, NAVY_L, rounded=True, radius=0.06)
    rect(s, x + 0.32, 3.45, 0.45, 0.45, ORANGE, rounded=True, radius=0.2)
    text(s, x + 0.32, 4.15, cw - 0.64, 1.5,
         [{"txt": h, "size": 13.5, "bold": True, "color": WHITE, "spacing": 1.1},
          {"txt": b, "size": 10, "color": WHITE_D, "space_before": 7, "spacing": 1.3}])

# =============================================================== slide 17 ===
s = slide(WHITE)
heading(s, "A closed-loop Smart School ecosystem", top=0.65, size=30)
rect(s, M, 2.15, 5.9, 3.4, GREY_BG, line=LINE, rounded=True, radius=0.07)
text(s, M + 0.4, 2.5, 5.1, 2.6,
     [{"txt": "Blended learning", "size": 17, "bold": True, "color": NAVY, "spacing": 1.05},
      {"txt": "Hybrid teaching works seamlessly, online and offline alike. A flexible plugin "
              "architecture allows customisation and extension for each specific subject.",
       "size": 11, "color": TEXT, "space_before": 8, "spacing": 1.35},
      {"txt": "School server  —  0 ms latency  —  classroom display",
       "size": 10.5, "bold": True, "color": BLUE, "space_before": 12, "spacing": 1.2}])
rect(s, 6.85, 2.15, 5.9, 3.4, GREY_BG, line=LINE, rounded=True, radius=0.07)
text(s, 7.25, 2.5, 5.1, 2.6,
     [{"txt": "A closed loop", "size": 17, "bold": True, "color": NAVY, "spacing": 1.05},
      {"txt": "Data moves automatically from lesson prep through to the final report, with nothing "
              "dropped along the way.",
       "size": 11, "color": TEXT, "space_before": 8, "spacing": 1.35},
      {"txt": "Lesson prep  →  Teaching in class  →  Data flows back to the LMS",
       "size": 10.5, "bold": True, "color": ORANGE, "space_before": 12, "spacing": 1.2}])

# =============================================================== slide 18 ===
s = slide(WHITE)
heading(s, "How it is rolled out", top=0.65, size=30)
steps = [("Step 1", NAVY_L,
          "Put in a high-performance wide-area network. Set up cloud storage built to "
          "information-security standards, so resolutions and documents sit behind a wall separate "
          "from the school's internal internet."),
         ("Step 2", ORANGE,
          "Set up the content management system (LCMS) and configure the AI models that analyse "
          "internal data. Adapt the teaching equipment as needed."),
         ("Step 3", NAVY,
          "Train staff and run a pilot at the school. Show people how to get the most out of the "
          "software and how to use AI when writing papers and clarifying an argument.")]
sw_ = (SW - 2 * M - 1.0) / 3
for i, (label, col, body) in enumerate(steps):
    x = M + i * (sw_ + 0.5)
    text(s, x, 2.3, sw_, 0.8, [{"txt": label, "size": 34, "bold": True, "color": col}])
    rect(s, x, 3.3, sw_, 0.02, LINE)
    text(s, x, 3.6, sw_, 2.0, [{"txt": body, "size": 11, "color": TEXT, "spacing": 1.4}])
    if i < 2:
        rect(s, x + sw_ + 0.1, 3.3, 0.3, 0.02, LINE)

# =============================================================== slide 19 ===
s = slide()
rect(s, M, 1.15, SW - 2 * M, 0.02, NAVY_L)
text(s, M, 2.85, 9.5, 1.8,
     [{"txt": "THANK YOU", "size": 60, "bold": True, "color": NAVY, "spacing": 1.05}])
rect(s, M, 6.35, SW - 2 * M, 0.02, NAVY_L)

out = "Smart School - EN deck.pptx"
prs.save(out)
print("saved:", out, "| slides:", len(prs.slides.__iter__.__self__._sldIdLst))
