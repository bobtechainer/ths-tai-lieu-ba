# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from copy import copy

# Open original file
src = openpyxl.load_workbook(r'c:\Trường học số - source code\Kịch bản Demo THS - Mobifone Solution.xlsx')
wb = openpyxl.Workbook()

# === COPY ORIGINAL SHEETS (P2 + Admin) with formatting ===
def copy_sheet(src_ws, dst_ws):
    for row in src_ws.iter_rows():
        for cell in row:
            new_cell = dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.fill = copy(cell.fill)
                new_cell.alignment = copy(cell.alignment)
                new_cell.border = copy(cell.border)
                new_cell.number_format = cell.number_format
    for mc in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(mc))
    for col_letter, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[col_letter].width = dim.width
    for row_num, dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[row_num].height = dim.height

# Copy P2 as first sheet
ws_p2 = wb.active
ws_p2.title = src.sheetnames[0]
copy_sheet(src[src.sheetnames[0]], ws_p2)

# Copy Admin as second sheet
ws_admin = wb.create_sheet(title=src.sheetnames[1])
copy_sheet(src[src.sheetnames[1]], ws_admin)

# === STYLING HELPERS ===
WHITE = "FFFFFF"
GRAY = "F2F2F2"
thin_b = Border(
    left=Side("thin", "BFBFBF"), right=Side("thin", "BFBFBF"),
    top=Side("thin", "BFBFBF"), bottom=Side("thin", "BFBFBF"))
wrap_top = Alignment(wrap_text=True, vertical="top")
wrap_center = Alignment(wrap_text=True, vertical="center", horizontal="center")
HEADERS = ["STT", "Vai trò", "Cửa sổ thao tác", "Tính năng trải nghiệm", "Thao tác chi tiết", "Kết quả mong đợi"]
COL_W = [6, 16, 16, 28, 60, 35]

def make_sheet(ws, title, subtitle, dark, light, rows):
    ws.sheet_properties.tabColor = dark
    for i, w in enumerate(COL_W, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # Title
    ws.merge_cells("A1:F1")
    c = ws.cell(1, 1, title)
    c.font = Font("Arial", 14, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=dark)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 35
    # Subtitle
    ws.merge_cells("A2:F2")
    c = ws.cell(2, 1, subtitle)
    c.font = Font("Arial", 11, italic=True, color="333333")
    c.fill = PatternFill("solid", fgColor=light)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 45
    # Headers
    hf = PatternFill("solid", fgColor=dark)
    for j, h in enumerate(HEADERS, 1):
        c = ws.cell(3, j, h)
        c.font = Font("Arial", 11, bold=True, color=WHITE)
        c.fill = hf; c.alignment = wrap_center; c.border = thin_b
    ws.row_dimensions[3].height = 28
    # Data
    for i, rd in enumerate(rows):
        r = 4 + i
        stripe = PatternFill("solid", fgColor=GRAY if i % 2 else WHITE)
        for j, val in enumerate(rd, 1):
            c = ws.cell(r, j, val)
            c.font = Font("Arial", 10, bold=(j == 4))
            c.alignment = wrap_center if j <= 3 else wrap_top
            c.fill = stripe; c.border = thin_b
        lines = max(str(v).count("\n") + 1 for v in rd if v)
        ws.row_dimensions[r].height = max(20, min(lines * 15, 220))
    ws.freeze_panes = "A4"

# === P3: GIAO BTVN ===
P3_ROWS = [
    [1, "Giáo viên", "Giáo viên", "GV tạo bài tập về nhà (nhiệm vụ) cho lớp",
     "B1: Vào lớp học (VD: lớp Hóa học 10A1) --> tab \"Nhiệm vụ\" --> bấm \"Tạo nhiệm vụ\"\n"
     "B2: Bước 1/3 - Thông tin chung:\n"
     "- Nhập Tiêu đề: \"Ôn tập Chương 6 - Tốc độ phản ứng\"\n"
     "- Chọn đối tượng: tick chọn lớp 10A1 (hệ thống hiển thị danh sách HS lớp, GV có thể chọn cả lớp hoặc chọn riêng từng HS)\n"
     "- Chọn Hạn nộp: VD 23:59 ngày mai\n"
     "- Tick \"Cho phép nộp muộn\" nếu muốn\n"
     "B3: Bước 2/3 - Nội dung bài tập:\n"
     "- GV tự tạo quiz trắc nghiệm ngay trong hệ thống: bấm \"Thêm câu hỏi\" --> chọn dạng (Chọn 1 đáp án / Nhiều đáp án / Đúng-Sai / Điền khuyết...)\n"
     "- VD tạo 5 câu trắc nghiệm về tốc độ phản ứng, nhập đề + 4 đáp án + chọn đáp án đúng\n"
     "- Hoặc giao dạng tự luận: HS upload file bài làm\n"
     "B4: Bước 3/3 - Xem lại và bấm \"Giao nhiệm vụ\"",
     "Nhiệm vụ xuất hiện trong danh sách lớp.\nHS nhận được thông báo bài tập mới trên hệ thống.\nURL: /lop-hoc/[classroomId] --> tab Nhiệm vụ"],

    [2, "Học sinh", "Học sinh", "HS nhận bài tập và làm bài trên hệ thống",
     "B1: HS đăng nhập --> vào mục \"Nhiệm vụ\" hoặc vào lớp Hóa học 10A1 --> tab Nhiệm vụ\n"
     "B2: Thấy danh sách nhiệm vụ được giao kèm deadline (VD: \"Ôn tập Chương 6 - Tốc độ phản ứng\" - Hạn nộp: 23:59 ngày mai)\n"
     "B3: Bấm vào nhiệm vụ --> bấm \"Làm bài\"\n"
     "B4: Trả lời lần lượt 5 câu trắc nghiệm (mỗi câu chọn 1 đáp án)\n"
     "- Hệ thống hiển thị thanh tiến độ: đã trả lời 3/5 câu\n"
     "B5: Bấm \"Nộp bài\" --> hệ thống tự chấm quiz ngay lập tức\n"
     "B6: Xem kết quả: 4/5 câu đúng = 8 điểm, hiển thị đáp án đúng/sai từng câu",
     "HS thấy rõ deadline, tiến độ nộp bài.\nHệ thống tự chấm quiz trắc nghiệm ngay sau khi nộp.\nHS biết ngay kết quả + đáp án đúng."],

    [3, "Giáo viên", "Giáo viên", "GV xem Dashboard chấm điểm và nhận xét bài nộp",
     "B1: GV vào chi tiết nhiệm vụ \"Ôn tập Chương 6\" đã giao\n"
     "B2: Hệ thống hiển thị Dashboard chấm điểm gồm:\n"
     "- Header: tên nhiệm vụ, lớp, deadline, trạng thái\n"
     "- Tổng quan: tỷ lệ hoàn thành (VD: 28/35 HS đã nộp = 80%), độ chính xác trung bình (VD: 72%)\n"
     "- Phân bố điểm theo 4 mức: 9-10 điểm (5 HS), 7-8 điểm (12 HS), 5-6 điểm (8 HS), dưới 5 (3 HS)\n"
     "- HS cao nhất: Nguyễn Văn A (10 điểm), HS thấp nhất: Trần Văn B (3 điểm)\n"
     "B3: Tab \"Học sinh\" --> bảng danh sách từng HS: tên, trạng thái (đã nộp/chưa nộp/nộp muộn), điểm, số câu đúng/sai\n"
     "B4: Click vào 1 HS --> trang chấm chi tiết:\n"
     "- Xem bài làm từng câu, đáp án HS chọn vs đáp án đúng\n"
     "- GV viết nhận xét cho HS (AI Tag gợi ý nhận xét tự động)\n"
     "- Bấm \"Hiển thị nhận xét\" để HS thấy được đánh giá của GV\n"
     "B5: Tab \"Danh sách câu hỏi\" --> phân tích từng câu: % HS trả lời đúng, phân bố đáp án --> GV biết câu nào HS yếu nhất\n"
     "URL: /nhiem-vu/[missionId]",
     "GV có dashboard chấm điểm trực quan.\nPhân bố điểm 4 mức + top/bottom HS.\nAI Tag gợi ý nhận xét tiết kiệm thời gian.\nPhân tích từng câu hỏi giúp GV biết HS yếu ở đâu."],

    [4, "Học sinh", "Học sinh", "HS xem kết quả và nhận xét từ GV",
     "B1: HS vào chi tiết nhiệm vụ đã nộp, xem được:\n"
     "- Điểm số: 8/10\n"
     "- Nhận xét của GV (nếu GV đã bật hiển thị)\n"
     "B2: Xem chi tiết từng câu đúng/sai, so sánh đáp án mình chọn vs đáp án đúng\n"
     "B3: Xem tiến độ hoàn thành: đã hoàn thành 3/5 nhiệm vụ trong tuần",
     "HS nắm được điểm mạnh/yếu của mình.\nCó hướng cải thiện cụ thể từ nhận xét GV."],
]

# === P4: BÁO CÁO ĐA BÊN ===
P4_ROWS = [
    [1, "Giáo viên", "Smart Class\n(GV)", "GV xem Dashboard thống kê tuần của mình",
     "B1: GV đăng nhập Smart Class --> vào Dashboard\n"
     "B2: Xem stat cards tuần:\n"
     "- Số tiết dạy: 12 tiết\n"
     "- Tổng giờ giảng dạy: 9.5 giờ\n"
     "- Tổng số HS: 140 HS\n"
     "- Tỷ lệ chuyên cần: 92%\n"
     "URL: smart-class.truonghocsoquocgia.vn --> Dashboard",
     "GV nắm tổng quan hoạt động giảng dạy tuần qua."],

    [2, "Giáo viên", "Smart Class\n(GV)", "GV xem báo cáo buổi học chi tiết (sau mỗi tiết dạy)",
     "B1: Kết thúc buổi dạy --> bấm \"Xem báo cáo\"\n"
     "B2: Biểu đồ Mức độ tập trung (EngagementChart): timeline cho thấy lúc nào HS tập trung cao/thấp\n"
     "B3: Kết quả quiz trong buổi (QuizResultsChart): % đúng từng câu\n"
     "B4: Bảng xếp hạng HS (LeaderboardChart): top HS theo accuracy + tốc độ trả lời\n"
     "B5: So sánh giữa các lớp (CompareDashboard): VD so sánh 10A1 vs 10A2 cùng bài giảng",
     "Dữ liệu real-time, biểu đồ trực quan.\nGV biết lúc nào HS mất tập trung, câu nào HS yếu."],

    [3, "Học sinh", "Smart Class\n(HS)", "HS xem báo cáo cá nhân sau buổi học",
     "B1: HS mở StudentClassReport sau khi lớp kết thúc\n"
     "B2: Xem điểm số, thứ hạng trong lớp, tỷ lệ trả lời đúng\n"
     "B3: Tab \"Nhận xét\" --> AI đánh giá cá nhân (aiAdvice):\n"
     "- Điểm mạnh: \"Em trả lời nhanh và chính xác ở phần tốc độ phản ứng\"\n"
     "- Điểm yếu: \"Em cần ôn lại phần yếu tố ảnh hưởng đến tốc độ phản ứng\"\n"
     "- Chiến lược: \"Nên làm thêm bài tập về xúc tác và nhiệt độ\"\n"
     "B4: Tab \"So sánh\" --> Radar chart so điểm cá nhân với trung bình lớp\n"
     "B5: Tab \"Mức độ tập trung\" --> biểu đồ engagement theo timeline buổi học",
     "HS biết đứng ở đâu so với lớp.\nAI gợi ý cải thiện cá nhân hóa."],

    [4, "Học sinh", "Smart Class\n(HS)", "HS xem thống kê tuần cá nhân",
     "B1: Xem Dashboard HS\n"
     "B2: Stat cards:\n"
     "- Số buổi tham gia: 8/10 buổi\n"
     "- Giờ học: 6.5 giờ\n"
     "- Tỷ lệ chuyên cần: 80%",
     "HS tự theo dõi được tần suất học tập mỗi tuần."],

    [5, "GV --> PHHS", "Site trường\n(GV)", "GV gửi phiếu báo cáo học tập cho Phụ huynh qua email",
     "B1: GV vào lớp Hóa học 10A1 --> tab \"Báo cáo học sinh\"\n"
     "URL: /lop-hoc/[classroomId]/bao-cao-hoc-sinh\n"
     "B2: Chọn HS cần gửi báo cáo (VD: Nguyễn Văn A), chọn tuần báo cáo (Tuần 4 tháng 7)\n"
     "B3: Hệ thống hiển thị Phiếu báo cáo học tập gồm:\n"
     "- Biểu đồ chuyên cần: đi học 4/5 buổi, vắng 1 buổi (có phép)\n"
     "- Phân bố điểm nhiệm vụ: 3 nhiệm vụ, điểm lần lượt 8, 9, 7\n"
     "- Tỷ lệ hoàn thành: 100% (3/3 nhiệm vụ)\n"
     "- Danh sách nhiệm vụ chi tiết: tên + điểm + trạng thái\n"
     "- Hoạt động học tập: lịch sử truy cập, thời gian online\n"
     "B4: GV bấm nút \"Gửi cho phụ huynh\" --> mở form soạn email:\n"
     "- Tiêu đề mẫu: \"Báo cáo học tập tuần 4/7 - Nguyễn Văn A\"\n"
     "- Nội dung mẫu: \"Kính gửi Quý phụ huynh, dưới đây là báo cáo tình hình học tập của con...\"\n"
     "- GV có thể tùy chỉnh nhận xét thêm\n"
     "B5: Bấm \"Gửi ngay\" --> hệ thống tự xuất PDF phiếu báo cáo, upload S3, gửi email đến parent_emails của HS\n"
     "Lưu ý: Email PHHS cần được nhập sẵn trong profile HS (mục \"Email phụ huynh\")",
     "PHHS nhận email kèm file PDF phiếu báo cáo học tập chi tiết.\nNắm được chuyên cần, điểm số, hoàn thành nhiệm vụ của con.\nGV tiết kiệm thời gian nhờ form mẫu sẵn + PDF tự động."],
]

# === BUILD NEW SHEETS ===
ws3 = wb.create_sheet(title="P3 GIAO BTVN - CHẤM ĐIỂM")
make_sheet(ws3,
    "Kịch bản 3 - Sau giờ học: Giao và làm bài tập về nhà, chấm điểm",
    "Mục tiêu: GV giao BTVN cho HS sau buổi học. HS nhận bài, làm bài, nộp trên hệ thống. GV chấm điểm kèm nhận xét. AI Tag hỗ trợ đánh giá. Data demo: Lớp Hóa học 10A1, bài \"Ôn tập Chương 6 - Tốc độ phản ứng\", 5 câu trắc nghiệm.",
    "BF6B00", "FFF2CC", P3_ROWS)

ws4 = wb.create_sheet(title="P4 BÁO CÁO GV-HS-PHHS")
make_sheet(ws4,
    "Kịch bản 4 - Theo dõi & phân tích: Báo cáo cho GV, HS, Phụ huynh",
    "Mục tiêu: Tổng hợp kết quả học tập cho 3 bên. GV nắm tình hình lớp + so sánh. HS tự theo dõi cá nhân + AI gợi ý. PHHS nhận báo cáo email PDF. Data demo: GV dạy 12 tiết/tuần, HS Nguyễn Văn A lớp 10A1, email PHHS: phuhuynhdemo@gmail.com.",
    "6B2FA0", "E8D5F5", P4_ROWS)

# === REORDER: P2 -> P3 -> P4 -> Admin ===
wb.move_sheet("P3 GIAO BTVN - CHẤM ĐIỂM", offset=-1)
wb.move_sheet("P4 BÁO CÁO GV-HS-PHHS", offset=-1)

out = r"c:\Trường học số - source code\Kịch bản Demo THS - Nâng cao.xlsx"
wb.save(out)
print("OK: " + out)
