import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# Colors
BLUE_DARK = "1B3A5C"
BLUE_MED = "2E75B6"
BLUE_LIGHT = "D6E4F0"
GREEN_DARK = "1D6B3F"
GREEN_LIGHT = "E2EFDA"
ORANGE_DARK = "BF6B00"
ORANGE_LIGHT = "FFF2CC"
PURPLE_DARK = "6B2FA0"
PURPLE_LIGHT = "E8D5F5"
RED_DARK = "A03030"
RED_LIGHT = "FCE4EC"
WHITE = "FFFFFF"
GRAY_BG = "F2F2F2"

# Styles
header_font = Font(name="Arial", size=11, bold=True, color=WHITE)
title_font = Font(name="Arial", size=14, bold=True, color=WHITE)
subtitle_font = Font(name="Arial", size=11, italic=True, color="333333")
cell_font = Font(name="Arial", size=10)
bold_font = Font(name="Arial", size=10, bold=True)
thin_border = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
wrap_align = Alignment(wrap_text=True, vertical="top")
center_align = Alignment(wrap_text=True, vertical="center", horizontal="center")

COL_WIDTHS = [6, 18, 18, 30, 55, 35]
HEADERS = ["STT", "Vai trò", "Cửa sổ thao tác", "Tính năng trải nghiệm", "Thao tác chi tiết", "Kết quả mong đợi"]

SHEETS = [
    {
        "name": "P1 TRƯỚC KHI LÊN LỚP",
        "title": "Kịch bản 1 - Trước khi lên lớp: Tự học trước/ tự chuẩn bị bài trước ở nhà",
        "subtitle": "Mục tiêu: Học sinh tự học trước ở nhà trước khi lên lớp. Giáo viên tham khảo học liệu số và chuẩn bị bài giảng tại nhà.",
        "color_dark": BLUE_DARK,
        "color_med": BLUE_MED,
        "color_light": BLUE_LIGHT,
        "rows": [
            [1, "GV / HS\n(chung)", "GV / HS",
             "Đọc sách giáo khoa phiên bản tương tác sinh động, ghi nhớ kiến thức tốt hơn",
             "B1: Đv tài khoản GV: Vào Quản trị đào tạo --> Kho học liệu số --> Bộ sách\nĐv tài khoản HS: Vào Học tập --> Kho học liệu số --> Bộ sách\nB2: Chọn sách Hóa học 10 - Chương 6: Tốc độ phản ứng - Bài 6\nB3: Đọc sách và nhờ AI Tutor hỗ trợ\n- Crop ảnh nội dung bất kì trong sách gửi AI Tutor giải đáp\n- Làm quiz sai --> Bấm Kiểm tra kết quả, AI Tutor đưa ra nhận xét và chiến lược giải bài\n- Làm quiz đúng --> Bấm Giải thích, AI giải đáp sâu hơn",
             "- Trải nghiệm sách phiên bản tương tác sinh động, mượt mà\n- AI Tutor nhận diện được ngữ cảnh trong sách và đưa ra giải đáp chính xác"],
            [2, "GV / HS\n(chung)", "GV / HS",
             "Tham khảo thêm học liệu khác tại Kho học liệu để ôn tập",
             "B1: Bấm chọn Kho học liệu tìm thêm video, mô hình, tài liệu về môn học\nB2: Tìm và click xem một học liệu bất kì\n- Mô hình mô phỏng 3D: xoay, zoom, tương tác như thực tế\n- Bộ đề để tự luyện: tự luyện đề và xem đáp án, AI hỗ trợ",
             "Học liệu load nhanh, xem ổn định, tương tác mô hình 3D mượt"],
            [3, "GV / HS\n(chung)", "GV / HS",
             "Hỏi đáp kiến thức học tập cùng AI",
             "B1: Mở chatbot AI Tutor ở góc dưới bên trái màn hình\nB2: Hỏi bất kì kiến thức nào để AI giải đáp\nB3: Gửi ảnh đề bài chữ viết tay để AI nhận diện và giải đáp, hoặc chụp ảnh bài tập gửi AI chấm",
             "- AI Tutor phản hồi chính xác, rõ ràng\n- Với bài phức tạp, AI đưa ra chiến lược giải từng bước thay vì chép đáp án trực tiếp"],
        ],
    },
    {
        "name": "P2 TRONG GIỜ HỌC",
        "title": "Kịch bản 2 - Trong giờ học: Lớp học thông minh",
        "subtitle": "Mục tiêu: Giáo viên tổ chức lớp học theo mô hình lớp học thông minh ngay trên hệ thống thay vì bảng đen truyền thống. Học sinh tương tác trực tiếp trên tablet/thiết bị cá nhân.",
        "color_dark": GREEN_DARK,
        "color_med": GREEN_DARK,
        "color_light": GREEN_LIGHT,
        "rows": [
            [1, "Giáo viên", "Giáo viên", "GV mở phòng học và tiến hành điểm danh",
             "B1: Vào Quản trị đào tạo --> Lớp học thông minh\nB2: GV mở lớp theo thời khóa biểu (thường trước 5 phút), hệ thống tự động sinh mã Passcode 6 số và QR Code\nB3: GV mở phiên điểm danh để HS quét mặt qua camera AI",
             "Phòng học mở thành công, hiển thị giao diện bảng trắng"],
            [2, "Học sinh", "Học sinh", "HS tham gia và tự điểm danh quét khuôn mặt",
             "B1: HS chọn đúng tiết học theo TKB để vào lớp\nB2: HS quét mặt qua camera AI để điểm danh",
             "HS vào phòng thành công. Bên màn hình GV hiển thị danh sách HS tham gia."],
            [3, "Giáo viên", "Giáo viên", "GV bắt đầu tổ chức giảng dạy",
             "B1: GV mở sách giáo khoa đã xem trước ở nhà, VD môn Hóa\n- Tại thanh công cụ phía dưới, bấm icon \"Hộp phép thuật\" --> chọn \"Sách\"\n- Chọn Hóa học lớp 10 - KNTT, Chương 6: Tốc độ phản ứng\nB2: GV viết lên bảng trắng trên hệ thống (bút viết hoặc gõ text)\nB3: GV hỏi đáp, thí nghiệm và tổ chức game quiz\n- Mở bảng tuần hoàn: \"Hộp phép thuật\" --> \"Tuần hoàn\"\n- Thí nghiệm hóa học: \"Hộp phép thuật\" --> \"Canvas app\" --> \"Thí nghiệm hóa học\" --> \"Kho thí nghiệm\" --> chọn thí nghiệm phù hợp\n- Game quiz: \"Hộp phép thuật\" --> \"Gamequiz\"",
             "GV tổ chức được lớp học tương tác thông minh, thay thế bảng đen truyền thống"],
            [4, "Học sinh", "Học sinh", "HS tương tác cùng GV trong buổi học",
             "HS theo dõi bảng trắng trên tablet/thiết bị cá nhân:\n- Xem GV mở sách, viết chữ lên bảng\n- Ghi chép hoặc record lại buổi dạy\n- Bấm giơ tay phát biểu hoặc thao tác trực tiếp lên bảng khi GV mở khóa\n- Xem bảng tuần hoàn GV mở\n- Thực hành thí nghiệm ngay trên tablet\n- Hỏi AI Tutor nếu ko hiểu: bấm Mũi tên --> chọn AI Tutor\n- Tham gia game quiz, trả lời nhanh để lọt top BXH",
             "HS theo dõi được nội dung GV dạy trên bảng, tương tác và tham gia cùng GV"],
            [5, "Giáo viên", "Giáo viên", "GV kết thúc buổi học và xem báo cáo",
             "B1: GV bấm Kết thúc --> bấm Xem báo cáo học sinh\nB2: Xem chi tiết báo cáo tình hình buổi học: mức độ tham gia, mức độ tập trung, điểm danh, nhận xét buổi học",
             "Dữ liệu cập nhật real-time chính xác. Phòng học đóng lại ở cả 2 bên."],
        ],
    },
    {
        "name": "P3 SAU GIỜ HỌC - BTVN",
        "title": "Kịch bản 3 - Sau giờ học: Giao và làm bài tập về nhà",
        "subtitle": "Mục tiêu: GV giao BTVN (nhiệm vụ) cho HS sau buổi học. HS nhận bài, làm bài và nộp trên hệ thống. GV chấm điểm, nhận xét. AI Tag hỗ trợ đánh giá.",
        "color_dark": ORANGE_DARK,
        "color_med": ORANGE_DARK,
        "color_light": ORANGE_LIGHT,
        "rows": [
            [1, "Giáo viên", "Giáo viên", "GV tạo bài tập về nhà (nhiệm vụ) cho lớp",
             "B1: Vào lớp học --> tab \"Nhiệm vụ\" --> bấm \"Tạo nhiệm vụ\"\nB2: Bước 1 - Chọn lớp nhận bài (có thể giao cho 1 hoặc nhiều lớp cùng lúc)\nB3: Bước 2 - Đặt tên nhiệm vụ, mô tả yêu cầu, hạn nộp bài\nB4: Bước 3 - Chọn dạng bài: GV tự tạo quiz trắc nghiệm ngay trong nhiệm vụ hoặc giao bài tự luận/upload file. GV có thể tạo bộ câu hỏi trắc nghiệm trực tiếp, nhập đáp án đúng để hệ thống tự chấm\nB5: Bấm \"Giao nhiệm vụ\"",
             "Nhiệm vụ xuất hiện trong danh sách lớp, HS nhận được thông báo bài tập mới"],
            [2, "Học sinh", "Học sinh", "HS nhận bài tập và làm bài",
             "B1: HS vào mục \"Nhiệm vụ\", thấy danh sách bài tập được giao kèm deadline\nB2: Chọn nhiệm vụ cần làm --> bấm \"Làm bài\"\nB3: Trả lời quiz trắc nghiệm / viết tự luận / upload file bài làm\nB4: Nộp bài",
             "HS thấy rõ deadline, tiến độ, trạng thái nộp/chưa nộp. Hệ thống tự chấm quiz trắc nghiệm ngay sau khi nộp."],
            [3, "Giáo viên", "Giáo viên", "GV xem tổng quan và chấm điểm bài nộp",
             "B1: GV vào chi tiết nhiệm vụ đã giao, hệ thống hiển thị Dashboard chấm điểm\nB2: Xem tổng quan: số HS đã nộp/chưa nộp, trạng thái deadline, phân bố điểm tổng quát\nB3: Xem danh sách bài nộp từng HS trong bảng, chấm từng bài\nB4: AI Tag tự động gợi ý nhận xét, GV có thể sử dụng hoặc tự viết nhận xét",
             "GV có dashboard chấm điểm rõ ràng, AI hỗ trợ tiết kiệm thời gian nhận xét"],
            [4, "Học sinh", "Học sinh", "HS xem kết quả và phản hồi từ GV",
             "B1: HS vào chi tiết nhiệm vụ đã nộp, xem được điểm số, nhận xét của GV\nB2: Xem tiến độ hoàn thành nhiệm vụ\nB3: Xem chi tiết từng câu hỏi đúng/sai, so sánh đáp án",
             "HS nắm được điểm mạnh/yếu của mình, có hướng cải thiện cụ thể"],
        ],
    },
    {
        "name": "P4 BÁO CÁO ĐA BÊN",
        "title": "Kịch bản 4 - Theo dõi và phân tích: Báo cáo cho GV, HS, PHHS",
        "subtitle": "Mục tiêu: Tổng hợp kết quả học tập và phản hồi cho cả 3 bên: GV nắm tình hình lớp, HS tự theo dõi tiến độ cá nhân, PHHS nhận báo cáo qua email kèm phiếu học tập chi tiết.",
        "color_dark": PURPLE_DARK,
        "color_med": PURPLE_DARK,
        "color_light": PURPLE_LIGHT,
        "rows": [
            [1, "Giáo viên", "Giáo viên", "GV xem thống kê tuần của mình",
             "B1: Vào Dashboard Smart Class\nB2: Xem các chỉ số: số tiết dạy trong tuần, tổng giờ giảng dạy, tổng số HS, tỷ lệ chuyên cần",
             "GV nắm tổng quan hoạt động giảng dạy của mình trong tuần"],
            [2, "Giáo viên", "Giáo viên", "GV xem báo cáo buổi học chi tiết",
             "B1: Kết thúc buổi dạy --> bấm \"Xem báo cáo\"\nB2: Xem biểu đồ mức độ tập trung của cả lớp, kết quả quiz\nB3: Xem bảng xếp hạng HS theo độ chính xác và tốc độ trả lời\nB4: So sánh giữa các lớp khác nhau mà GV dạy",
             "Dữ liệu real-time, biểu đồ trực quan giúp GV đánh giá nhanh chất lượng buổi dạy"],
            [3, "Học sinh", "Học sinh", "HS xem báo cáo cá nhân sau buổi học",
             "B1: HS xem báo cáo cá nhân sau khi lớp kết thúc\nB2: Xem điểm số, thứ hạng trong lớp, tỷ lệ trả lời đúng\nB3: Tab \"Nhận xét\" có AI đánh giá cá nhân: điểm mạnh, điểm yếu, chiến lược cải thiện\nB4: Tab \"So sánh\" có radar chart so với trung bình lớp\nB5: Tab \"Mức độ tập trung\" hiển thị biểu đồ engagement theo timeline buổi học",
             "HS biết mình đang đứng ở đâu so với lớp, nhận gợi ý cải thiện từ AI"],
            [4, "Học sinh", "Học sinh", "HS xem thống kê tuần",
             "B1: Xem Dashboard HS\nB2: Xem các chỉ số: số buổi tham gia, giờ học, tỷ lệ chuyên cần trong tuần",
             "HS tự theo dõi được tần suất học tập của mình"],
            [5, "GV --> PHHS", "Giáo viên", "GV gửi phiếu báo cáo học tập cho Phụ huynh qua email",
             "B1: GV vào lớp --> tab \"Báo cáo học sinh\"\nB2: Chọn HS cần gửi báo cáo, chọn tuần cần báo cáo\nB3: Hệ thống hiển thị Phiếu báo cáo học tập gồm: biểu đồ chuyên cần, phân bố điểm nhiệm vụ, tỷ lệ hoàn thành, danh sách nhiệm vụ chi tiết\nB4: GV bấm \"Gửi cho phụ huynh\" --> mở form soạn email với tiêu đề và nội dung mẫu sẵn (\"Kính gửi Quý phụ huynh...\")\nB5: GV tùy chỉnh nhận xét --> bấm \"Gửi ngay\". Hệ thống tự đính kèm file PDF phiếu báo cáo và gửi email đến email phụ huynh đã đăng ký",
             "Phụ huynh nhận email kèm file PDF phiếu báo cáo học tập chi tiết, nắm được tình hình học tập của con"],
        ],
    },
    {
        "name": "P5 ADMIN TRƯỜNG",
        "title": "Kịch bản 5 - Admin trường: Dashboard và báo cáo nâng cao",
        "subtitle": "Mục tiêu: Quản trị viên nhà trường nắm tổng quan tình hình toàn trường, theo dõi báo cáo theo lớp/giáo viên, quản lý thi kiểm tra và liveclass.",
        "color_dark": RED_DARK,
        "color_med": RED_DARK,
        "color_light": RED_LIGHT,
        "rows": [
            [1, "Admin trường", "Admin trường", "Dashboard tổng quan trường",
             "B1: Đăng nhập tài khoản admin\nB2: Chọn tab Dashboard tại menu trái\nB3: Xem tỷ lệ tương tác SGK theo môn (biểu đồ tròn)\nB4: Xem thống kê truy cập hệ thống\nB5: Xem danh sách lớp cần chú ý (cảnh báo lớp có vấn đề)\nB6: Xem bảng xếp hạng lớp và hoạt động gần đây",
             "Nắm nhanh tình hình toàn trường, phát hiện sớm lớp cần hỗ trợ"],
            [2, "Admin trường", "Admin trường", "Báo cáo tổng quan chi tiết",
             "B1: Chọn \"Báo cáo tổng quan\"\nB2: Lọc theo tuần/tháng/khoảng thời gian tùy ý\nB3: Xem các chỉ số tổng hợp, biểu đồ tròn phân bố, biểu đồ đường xu hướng\nB4: Xem thống kê tương tác SGK",
             "Phân tích xu hướng sử dụng hệ thống theo thời gian"],
            [3, "Admin trường", "Admin trường", "Báo cáo theo lớp (tập trung BTVN)",
             "B1: Chọn tab Báo cáo theo lớp\nB2: Tìm kiếm hoặc lọc lớp muốn xem\nB3: Click vào chi tiết lớp\nB4: Tab Phân bố điểm: lọc theo nhiệm vụ, xem các mức điểm\nB5: Tab Danh sách nhiệm vụ: tên, trạng thái, ngày tạo, ngày hoàn thành, số HS hoàn thành vs tổng",
             "Xem được báo cáo từng lớp, đánh giá chất lượng BTVN và tỷ lệ hoàn thành"],
            [4, "Admin trường", "Admin trường", "Báo cáo theo giáo viên",
             "B1: Chọn Báo cáo theo giáo viên\nB2: Tìm kiếm GV muốn xem, click xem chi tiết\nB3: Tab \"Lớp giảng dạy\": danh sách lớp bộ môn + lớp chủ nhiệm\nB4: Tab \"Thống kê học liệu\": tổng số học liệu GV đã tạo",
             "Xem được báo cáo từng GV, số lớp dạy và học liệu số đã tạo"],
            [5, "Admin trường", "Admin trường", "Quản lý thi kiểm tra",
             "B1: Chọn tab Quản lý thi - kiểm tra\nB2: Tạo kỳ thi mới: click \"Thêm mới\", điền thông tin, lưu\nB3: Theo dõi kỳ thi: số thí sinh đạt/không đạt\nB4: Xem báo cáo bài thi chi tiết",
             "Theo dõi được toàn bộ hoạt động thi kiểm tra của trường"],
            [6, "Admin trường", "Admin trường", "Quản lý Liveclass",
             "B1: Chọn tab Live class\nB2: Tìm kiếm lớp theo thời gian\nB3: Chọn lớp bất kỳ, click icon chi tiết xem lịch sử buổi học\nB4: Danh sách HS tham gia, vắng/có mặt, chi tiết thời gian ra vào phòng",
             "Xem được danh sách buổi học, chi tiết từng HS từng buổi"],
        ],
    },
]


def build_sheet(ws, data):
    ws.sheet_properties.tabColor = data["color_dark"]
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    title_cell = ws.cell(row=1, column=1, value=data["title"])
    title_cell.font = title_font
    title_cell.fill = PatternFill("solid", fgColor=data["color_dark"])
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 35

    # Subtitle row
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
    sub_cell = ws.cell(row=2, column=1, value=data["subtitle"])
    sub_cell.font = subtitle_font
    sub_cell.fill = PatternFill("solid", fgColor=data["color_light"])
    sub_cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 40

    # Header row
    header_fill = PatternFill("solid", fgColor=data["color_med"])
    for j, h in enumerate(HEADERS, 1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = thin_border
    ws.row_dimensions[3].height = 28

    # Data rows
    for i, row_data in enumerate(data["rows"]):
        r = 4 + i
        stripe = PatternFill("solid", fgColor=GRAY_BG) if i % 2 == 1 else PatternFill("solid", fgColor=WHITE)
        for j, val in enumerate(row_data, 1):
            c = ws.cell(row=r, column=j, value=val)
            c.font = bold_font if j == 4 else cell_font
            c.alignment = center_align if j <= 3 else wrap_align
            c.fill = stripe
            c.border = thin_border
        # Auto row height based on content
        max_lines = max(str(v).count("\n") + 1 for v in row_data if v)
        ws.row_dimensions[r].height = max(20, min(max_lines * 15, 200))

    # Freeze header
    ws.freeze_panes = "A4"


# Build all sheets
for i, sheet_data in enumerate(SHEETS):
    if i == 0:
        ws = wb.active
        ws.title = sheet_data["name"]
    else:
        ws = wb.create_sheet(title=sheet_data["name"])
    build_sheet(ws, sheet_data)

output = r"c:\Trường học số - source code\Kịch bản Demo THS - Nâng cao.xlsx"
wb.save(output)
print(f"Saved to {output}")
