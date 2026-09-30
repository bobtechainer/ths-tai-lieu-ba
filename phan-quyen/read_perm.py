import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

wb = openpyxl.load_workbook(os.path.join(os.path.dirname(__file__), '(VNTIT - MBF) PHÂN QUYỀN TÍNH NĂNG CỦA SLL ĐIỆN TỬ.xlsx'))
print('ALL SHEETS:', wb.sheetnames)
print('---')
for name in wb.sheetnames:
    ws = wb[name]
    print(f'\n{"="*80}')
    print(f'SHEET: {name} (Rows: {ws.max_row}, Cols: {ws.max_column})')
    print(f'{"="*80}')
    
    if ws.merged_cells.ranges:
        print(f'Merged cells: {[str(m) for m in ws.merged_cells.ranges][:20]}')
    
    for idx, row_data in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 300), values_only=True), 1):
        if all(v is None for v in row_data):
            continue
        vals = list(row_data)
        while vals and vals[-1] is None:
            vals.pop()
        print(f'  R{idx}: {vals}')
