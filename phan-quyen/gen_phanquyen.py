import sys, re
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from copy import copy
from pq_data import D

perm_path = r'c:\Trường học số - source code\(VNTIT - MBF) PHÂN QUYỀN TÍNH NĂNG CỦA SLL ĐIỆN TỬ.xlsx'
wb = openpyxl.load_workbook(perm_path)

def extract_id(text):
    if not text: return None
    s = re.sub(r'(\d)\s+\.', r'\1.', str(text).strip())
    m = re.match(r'^(\d+(?:\.\d+)*)', s)
    return m.group(1) if m else None

def find(key):
    if not key: return ('','','')
    if key in D: return D[key]
    parts = key.split('.')
    for i in range(len(parts)-1, 0, -1):
        p = '.'.join(parts[:i])
        if p in D: return D[p]
    return ('','','')

# Copy existing cell styles from row 2 col 1 for header, row 3 col 1 for data
for sn in ['Phân quyền THCSTHPT', 'Phân quyền Tiểu học']:
    ws = wb[sn]
    
    # Read header style from row 2, col 1
    ref_h = ws.cell(row=2, column=1)
    h_font = copy(ref_h.font) if ref_h.font else Font(bold=True, size=10)
    h_fill = copy(ref_h.fill) if ref_h.fill and ref_h.fill.fgColor and ref_h.fill.fgColor.rgb and ref_h.fill.fgColor.rgb != '00000000' else PatternFill("solid", fgColor="4472C4")
    h_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    h_border = copy(ref_h.border) if ref_h.border else Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
    
    # Ensure header font is white bold
    h_font_new = Font(bold=True, color="FFFFFF", size=h_font.size or 10, name=h_font.name or 'Arial')
    
    # Read data style from row 3, col 1
    ref_d = ws.cell(row=3, column=1)
    d_font = copy(ref_d.font) if ref_d.font else Font(size=10)
    d_border = copy(ref_d.border) if ref_d.border else Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
    d_alignment = Alignment(wrap_text=True, vertical='top')
    
    # Write headers
    headers = ['Đơn vị PIC', 'Vị trí menu trên THS Admin', 'Yêu cầu chi tiết cho Dev']
    for i, h in enumerate(headers):
        c = ws.cell(row=2, column=4+i, value=h)
        c.font = h_font_new
        c.fill = h_fill
        c.alignment = h_alignment
        c.border = h_border
    
    # Also style row 1 merged title area - extend merge if needed
    # Just add border to new header cells in row 1
    for i in range(3):
        c1 = ws.cell(row=1, column=4+i)
        c1.border = h_border
    
    cur_id = None
    for r in range(3, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        b = ws.cell(row=r, column=2).value
        c_val = ws.cell(row=r, column=3).value
        if a is None and b is None and c_val is None:
            continue
        
        eid = extract_id(a)
        if eid: cur_id = eid
        
        pic, menu, dev = find(cur_id)
        
        # Check if this is a module header row (col A has value, col B is None)
        is_header = (a is not None and b is None and c_val is None)
        
        # Copy font style from col A of same row for consistency
        ref_cell = ws.cell(row=r, column=1)
        row_font = copy(ref_cell.font) if ref_cell.font else d_font
        row_fill = copy(ref_cell.fill) if ref_cell.fill and ref_cell.fill.fgColor and ref_cell.fill.fgColor.rgb and ref_cell.fill.fgColor.rgb != '00000000' else None
        row_border = copy(ref_cell.border) if ref_cell.border else d_border
        
        vals = [pic, menu, dev]
        for i, v in enumerate(vals):
            cell = ws.cell(row=r, column=4+i, value=v)
            cell.font = copy(row_font)
            cell.alignment = d_alignment
            cell.border = row_border
            if row_fill:
                cell.fill = copy(row_fill)
    
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 42
    ws.column_dimensions['F'].width = 58

out = r'c:\Trường học số - source code\YEU_CAU_PHAN_QUYEN_v3.xlsx'
wb.save(out)
print(f'Done! {out}')
